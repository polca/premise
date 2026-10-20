"""Replace the whole-ICE mechanical proxy in BEV/FCEV cars, without Brightway.

Run after refreshing lci-pass_cars.xlsx. Original non-target activities,
driving parameters and separate electrical modules are preserved. The final
energy workbook is patched as ZIP/XML, preserving its formula caches.
"""

from collections import defaultdict
from copy import deepcopy
import json
import math
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook

from premise.vehicle_energy_use import COMMON_AUXILIARY_BATTERY, SOURCE_KEY

ROOT = Path(__file__).resolve().parents[1]
INVENTORIES = ROOT / "premise/data/additional_inventories"
SOURCE_PATH = ROOT / "premise/data/transport/electric_mechanical_powertrain.json"
OLD = "market for internal combustion engine, passenger car"
PRODUCT = "mechanical powertrain, electric passenger car"
MARKET = f"market for {PRODUCT}"
AUXILIARY = f"market for {COMMON_AUXILIARY_BATTERY}"
USED = "used mechanical powertrain, electric passenger car"
NOTE = (
    "Electric mechanical-powertrain correction (2026-10-17): gearbox and cooling "
    "from Notter 2010 SI S3; motor/controls/charger remain separate. Historical "
    "auxiliary battery is retained as a provisional common component in the full "
    "vehicle, excluded from per-MJ energy equipment. Nominal component mass is "
    "partitioned, not resized. Modern gearbox design and auxiliary-battery "
    "specification/life remain unvalidated. https://doi.org/10.1021/es903729a"
)


def source_data():
    return json.loads(SOURCE_PATH.read_text())


def fractions():
    data = source_data()["historical_auxiliary"]
    auxiliary = data["retained_battery_kg"] / data["ice_reference_kg"]
    return 1 - auxiliary, auxiliary


def exchange(name, product, amount, location="GLO", unit="kilogram", comment=""):
    return {
        "name": name,
        "reference product": product,
        "amount": amount,
        "location": location,
        "unit": unit,
        "type": "technosphere",
        "comment": comment,
    }


def activity(name, product, exchanges, comment, production=1):
    return {
        "name": name,
        "reference product": product,
        "location": "GLO",
        "unit": "kilogram",
        "type": "processwithreferenceproduct",
        "production amount": production,
        "comment": comment,
        "exchanges": [
            {**exchange(name, product, production), "type": "production"},
            *exchanges,
        ],
    }


def inventories():
    data = source_data()
    mechanical = data["mechanical_reference"]
    reference = mechanical["retained_mass_kg"]
    retained, gross = defaultdict(float), defaultdict(float)
    for row in mechanical["components"]:
        retained[row["material"]] += row["retained_kg"] / reference
        gross[row["material"]] += row["retained_kg"] * row["gross_factor"] / reference
    assert math.isclose(sum(retained.values()), 1)
    feeds = [
        exchange(
            f"market for {material}",
            material,
            amount,
            "RoW" if material == "ethylene glycol" else "GLO",
            comment="Notter 2010 S3 gearbox/cooling only: retained mass * waste factor / 25 kg reference.",
        )
        for material, amount in gross.items()
    ]
    whole_mass = mechanical["whole_electric_reference_mass_kg"]
    for name, product, key, loc, unit in [
        (
            "market group for electricity, medium voltage",
            "electricity, medium voltage",
            "whole_electric_electricity_kWh",
            "GLO",
            "kilowatt hour",
        ),
        (
            "market group for heat, district or industrial, natural gas",
            "heat, district or industrial, natural gas",
            "whole_electric_heat_MJ",
            "GLO",
            "megajoule",
        ),
        (
            "heat production, light fuel oil, at industrial furnace 1MW",
            "heat, district or industrial, other than natural gas",
            "whole_electric_light_fuel_oil_MJ",
            "RoW",
            "megajoule",
        ),
        (
            "market group for tap water",
            "tap water",
            "whole_electric_water_kg",
            "GLO",
            "kilogram",
        ),
        (
            "market for transport, freight, lorry 16-32 metric ton, EURO6",
            "transport, freight, lorry 16-32 metric ton, EURO6",
            "whole_electric_road_freight_tkm",
            "RER",
            "ton kilometer",
        ),
        (
            "market for transport, freight train",
            "transport, freight train",
            "whole_electric_rail_freight_tkm",
            "Europe without Switzerland",
            "ton kilometer",
        ),
        (
            "market for road vehicle factory",
            "road vehicle factory",
            "whole_electric_vehicle_factory_unit",
            "GLO",
            "unit",
        ),
    ]:
        feeds.append(
            exchange(
                name,
                product,
                mechanical[key] / whole_mass,
                loc,
                unit,
                "Mass allocation of whole-electric-reference assembly activity, not a measured gearbox-only process.",
            )
        )
    iron = ["steel, low-alloyed, hot rolled", "reinforcing steel"]
    scrap_iron = sum(gross[n] - retained[n] for n in iron)
    scrap_al = gross["aluminium, cast alloy"] - retained["aluminium, cast alloy"]
    plastic_names = [
        "polyethylene, high density, granulate",
        "polyphenylene sulfide",
        "synthetic rubber",
    ]
    scrap_plastic = sum(gross[n] - retained[n] for n in plastic_names)
    feeds.extend(
        [
            exchange(
                "market for iron scrap, unsorted",
                "iron scrap, unsorted",
                -scrap_iron,
            ),
            exchange(
                "market for aluminium scrap, new",
                "aluminium scrap, new",
                -scrap_al,
                "RoW",
                comment="Native scrap market; RoW is the non-European destination for this global equipment reference.",
            ),
            exchange(
                "market for waste plastic, mixture",
                "waste plastic, mixture",
                -scrap_plastic,
                "RoW",
            ),
            exchange(f"market for {USED}", USED, -1),
        ]
    )
    manufacture = activity(
        "mechanical powertrain production, electric passenger car",
        PRODUCT,
        feeds,
        NOTE + " Retained gearbox/cooling mass closes to 1 kg; gross material "
        "purchases include source waste factors. Manufacturing auxiliaries are "
        "mass-allocated historical proxies. No combustion catalyst, starter, "
        "fuel system, traction motor, inverter, charger or traction battery is copied.",
    )
    market = activity(
        MARKET,
        PRODUCT,
        [exchange(manufacture["name"], PRODUCT, 1)],
        "One kg complete mechanical gearbox/cooling equipment per kg sold. " + NOTE,
    )
    treatment = activity(
        "treatment of used mechanical powertrain, electric passenger car",
        USED,
        [
            exchange(
                "market for iron scrap, unsorted",
                "iron scrap, unsorted",
                -sum(retained[n] for n in iron),
            ),
            exchange(
                "market for aluminium scrap, post-consumer",
                "aluminium scrap, post-consumer",
                -retained["aluminium, cast alloy"],
            ),
            exchange(
                "market for waste plastic, mixture",
                "waste plastic, mixture",
                -sum(retained[n] for n in plastic_names),
                "RoW",
            ),
            exchange(
                "market for hazardous waste, for incineration",
                "hazardous waste, for incineration",
                -retained["ethylene glycol"],
                "RoW",
                comment="Coolant disposal proxy; not a validated coolant-recovery process.",
            ),
            exchange(
                "market for manual treatment facility, waste electric and electronic equipment",
                "manual treatment facility, waste electric and electronic equipment",
                1.6e-8,
                unit="unit",
                comment="Generic facility-rate proxy from existing manual drivetrain treatment; no electronic/motor scrap fractions copied.",
            ),
        ],
        "Dedicated mechanical treatment: retained steel/aluminium, polymers/rubber "
        "and coolant sum to 1 kg. Uses native scrap markets and "
        "documented generic treatment proxies, not future recovery efficiencies.",
        production=-1,
    )
    used_market = activity(
        f"market for {USED}",
        USED,
        [exchange(treatment["name"], USED, -1)],
        "One kg used mechanical equipment to its dedicated treatment.",
        production=-1,
    )
    aux = data["historical_auxiliary"]
    mass = aux["retained_battery_kg"]
    auxiliary_production = activity(
        "auxiliary battery production, passenger car, historical proxy",
        COMMON_AUXILIARY_BATTERY,
        [
            exchange("market for lead", "lead", aux["retained_lead_kg"] / mass),
            exchange(
                "market for polyphenylene sulfide",
                "polyphenylene sulfide",
                aux["retained_case_proxy_PPS_kg"] * aux["case_gross_factor"] / mass,
            ),
            exchange(
                "market for sulfuric acid",
                "sulfuric acid",
                aux["retained_acid_kg"] / mass,
                "RoW",
            ),
            exchange(
                "market for waste plastic, mixture",
                "waste plastic, mixture",
                -aux["retained_case_proxy_PPS_kg"]
                * (aux["case_gross_factor"] - 1)
                / mass,
                "RoW",
            ),
            exchange(
                "market for scrap lead acid battery", "scrap lead acid battery", -1
            ),
        ],
        "Provisional Notter S2 starter-battery material proxy, not an identified "
        "modern EV battery: retained 13 kg lead, 4 kg case proxy and 1 kg acid "
        "per historical 18 kg reference. Preserves prior full-vehicle direct lead "
        "when assigned 18/275 of the old ICE proxy mass. Case chemistry, assembly, "
        "voltage/capacity, sizing, calendar replacements and treatment remain "
        "unvalidated. Common equipment excluded from per-MJ energy inventory.",
    )
    auxiliary_market = activity(
        AUXILIARY,
        COMMON_AUXILIARY_BATTERY,
        [exchange(auxiliary_production["name"], COMMON_AUXILIARY_BATTERY, 1)],
        auxiliary_production["comment"],
    )
    return [
        manufacture,
        market,
        treatment,
        used_market,
        auxiliary_production,
        auxiliary_market,
    ]


def write_inventory():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Electric mechanical equipment"
    sheet.append(["database", "premise electric mechanical equipment"])
    sheet.append([])
    columns = [
        "name",
        "amount",
        "location",
        "unit",
        "categories",
        "type",
        "reference product",
        "comment",
    ]
    for ds in inventories():
        sheet.append(["Activity", ds["name"]])
        for key, value in ds.items():
            if key not in {"name", "exchanges"}:
                sheet.append([key, value])
        sheet.append(["Exchanges"])
        sheet.append(columns)
        for e in ds["exchanges"]:
            sheet.append([e.get(column) for column in columns])
        sheet.append([])
    workbook.save(INVENTORIES / "lci-electric-car-mechanical-powertrain.xlsx")


def blocks(sheet):
    starts = [row[0].row for row in sheet if row[0].value == "Activity"]
    for start, end in zip(starts, starts[1:] + [sheet.max_row + 2]):
        header = (
            next(
                r
                for r in range(start + 1, end)
                if sheet.cell(r, 1).value == "Exchanges"
            )
            + 1
        )
        columns = {
            str(sheet.cell(header, c).value): c
            for c in range(1, sheet.max_column + 1)
            if sheet.cell(header, c).value is not None
        }
        yield start, end, header, columns


def rescaled_values(sheet, row, columns, factor):
    values = [sheet.cell(row, c).value for c in range(1, sheet.max_column + 1)]
    for field in ["amount", "loc", "minimum", "maximum"]:
        if field in columns and values[columns[field] - 1] is not None:
            values[columns[field] - 1] = float(values[columns[field] - 1]) * factor
    return values


def patch_cars(path):
    workbook = load_workbook(path)
    mechanical, auxiliary = fractions()
    changes = 0
    for sheet in workbook:
        for start, end, header, columns in reversed(list(blocks(sheet))):
            name = sheet.cell(start, 2).value
            if not name.startswith(
                (
                    "Passenger car, battery electric,",
                    "Passenger car, fuel cell electric,",
                )
            ):
                continue
            old = [
                r
                for r in range(header + 1, end - 1)
                if sheet.cell(r, columns["name"]).value == OLD
            ]
            if not old:
                if any(
                    sheet.cell(r, columns["name"]).value == MARKET
                    for r in range(header + 1, end - 1)
                ):
                    continue
                raise ValueError(f"Missing mechanical-powertrain source: {name}")
            (row,) = old
            original = float(sheet.cell(row, columns["amount"]).value)
            replacement = rescaled_values(sheet, row, columns, mechanical)
            common = rescaled_values(sheet, row, columns, auxiliary)
            for values, supplier, product in [
                (replacement, MARKET, PRODUCT),
                (common, AUXILIARY, COMMON_AUXILIARY_BATTERY),
            ]:
                values[columns["name"] - 1] = supplier
                values[columns["reference product"] - 1] = product
                values[columns["comment"] - 1] = NOTE
            for c, value in enumerate(replacement, 1):
                sheet.cell(row, c).value = value
            sheet.insert_rows(row + 1)
            for c, value in enumerate(common, 1):
                sheet.cell(row + 1, c).value = value
            for r in range(start + 1, header):
                key = sheet.cell(r, 1).value
                if key == "comment":
                    sheet.cell(r, 2).value = (
                        str(sheet.cell(r, 2).value or "") + " " + NOTE
                    )
                elif key == "vehicle inventory provenance":
                    record = json.loads(sheet.cell(r, 2).value)
                    record["electric mechanical powertrain correction"] = {
                        "source": "Notter 2010 SI S2/S3; separate electric gearbox/cooling and provisional common auxiliary battery",
                        "original proxy mass kg": original,
                        "mechanical mass kg": original * mechanical,
                        "historical auxiliary mass kg": original * auxiliary,
                        "boundary": NOTE,
                    }
                    sheet.cell(r, 2).value = json.dumps(record)
            changes += 1
    if changes not in {0, 17}:
        raise ValueError(
            f"Expected 17 native BEV/FCEV cars or already-corrected inventory, found {changes}"
        )
    if changes:
        workbook.save(path)
    return changes


def patch_xml_cells(path, updates):
    ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    ET.register_namespace("", ns)
    with ZipFile(path) as archive:
        entries = [
            (entry, archive.read(entry.filename)) for entry in archive.infolist()
        ]
    for part, cells in updates.items():
        index = next(
            i for i, (entry, _) in enumerate(entries) if entry.filename == part
        )
        entry, content = entries[index]
        root = ET.fromstring(content)
        by_address = {cell.get("r"): cell for cell in root.iter(f"{{{ns}}}c")}
        for address, value in cells.items():
            cell = by_address[address]
            if cell.find(f"{{{ns}}}f") is not None:
                raise ValueError(f"Refusing to overwrite formula {part}:{address}")
            for child in list(cell):
                cell.remove(child)
            if isinstance(value, (float, int)):
                cell.attrib.pop("t", None)
                ET.SubElement(cell, f"{{{ns}}}v").text = repr(value)
            else:
                cell.set("t", "inlineStr")
                ET.SubElement(ET.SubElement(cell, f"{{{ns}}}is"), f"{{{ns}}}t").text = (
                    value
                )
        entries[index] = (entry, ET.tostring(root, encoding="utf-8"))
    with ZipFile(path, "w") as archive:
        for entry, content in entries:
            archive.writestr(entry, content)


def patch_energy(path):
    workbook = load_workbook(path)
    mechanical, auxiliary = fractions()
    updates = {}
    count = 0
    for index, sheet in enumerate(workbook, 1):
        cells = {}
        for start, end, header, columns in blocks(sheet):
            name = sheet.cell(start, 2).value
            if name not in {
                "electricity, used in passenger car",
                "hydrogen, used in a fuel cell in a passenger car",
            }:
                continue
            rows = [
                r
                for r in range(header + 1, end - 1)
                if sheet.cell(r, columns["name"]).value == OLD
            ]
            if not rows:
                if any(
                    sheet.cell(r, columns["name"]).value == MARKET
                    for r in range(header + 1, end - 1)
                ):
                    continue
                raise ValueError(f"Missing final-energy source proxy: {name}")
            (row,) = rows
            for field, value in [
                ("name", MARKET),
                ("reference product", PRODUCT),
                (
                    "amount",
                    float(sheet.cell(row, columns["amount"]).value) * mechanical,
                ),
            ]:
                cells[sheet.cell(row, columns[field]).coordinate] = value
            for r in range(start + 1, header):
                key = sheet.cell(r, 1).value
                if key == SOURCE_KEY:
                    record = json.loads(sheet.cell(r, 2).value)
                    for e in record["retained equipment"]:
                        if e["supplier"][0] == OLD:
                            common = deepcopy(e)
                            common["supplier"] = [
                                AUXILIARY,
                                COMMON_AUXILIARY_BATTERY,
                                "GLO",
                                "kilogram",
                            ]
                            common["amount per vehicle"] *= auxiliary
                            common.pop("amount per MJ", None)
                            common.pop("coefficient limitation", None)
                            common["boundary reason"] = (
                                "Provisional common auxiliary battery retained in full vehicle; excluded from energy-specific equipment."
                            )
                            record["excluded"].append(common)
                            e["supplier"] = [MARKET, PRODUCT, "GLO", "kilogram"]
                            e["amount per vehicle"] *= mechanical
                            e["amount per MJ"] *= mechanical
                            e["coefficient limitation"] = NOTE
                    record["electric mechanical powertrain correction"] = NOTE
                    cells[sheet.cell(r, 2).coordinate] = json.dumps(record)
                elif key == "comment":
                    cells[sheet.cell(r, 2).coordinate] = (
                        str(sheet.cell(r, 2).value or "") + " " + NOTE
                    )
            count += 1
        if cells:
            updates[f"xl/worksheets/sheet{index}.xml"] = cells
    if count not in {0, 2}:
        raise ValueError(f"Expected two electric-car energy wrappers, found {count}")
    if updates:
        patch_xml_cells(path, updates)
    return count


if __name__ == "__main__":
    write_inventory()
    print(
        "Electric car production activities corrected:",
        patch_cars(INVENTORIES / "lci-pass_cars.xlsx"),
    )
    print(
        "Final-energy wrappers corrected:",
        patch_energy(INVENTORIES / "lci-final-energy.xlsx"),
    )
