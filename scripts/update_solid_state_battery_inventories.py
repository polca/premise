"""Rebuild the documented baseline ASSB inventories from packaged source data.

Run with the premise environment from the repository root. This writes only
the dedicated ASSB workbook and three zero-share placeholders per mobile
scenario in the capacity workbook. It never reads or writes a Brightway DB.
"""

from copy import deepcopy
import json
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from bw2io import ExcelImporter
from openpyxl import Workbook

from premise.inventory_imports import migrate_import_db

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "premise/data"
INVENTORIES = DATA / "additional_inventories"
SOURCE = DATA / "battery/solid_state_reference_designs.json"
CHEMISTRIES = ("polymer", "oxidic", "sulfidic")
CELL_FRACTION = 0.72591


def technosphere(name, product, amount, location="GLO", unit="kilogram", comment=""):
    return dict(
        name=name,
        **{"reference product": product},
        amount=amount,
        location=location,
        unit=unit,
        type="technosphere",
        comment=comment,
    )


def biosphere(name, amount, categories):
    return dict(
        name=name,
        amount=amount,
        unit="kilogram",
        type="biosphere",
        categories=categories,
    )


def activity(name, product, exchanges, comment, unit="kilogram", **metadata):
    return {
        "name": name,
        "reference product": product,
        "location": "GLO",
        "unit": unit,
        "production amount": 1,
        "type": "processwithreferenceproduct",
        "comment": comment,
        **metadata,
        "exchanges": [
            dict(
                name=name,
                **{"reference product": product},
                amount=1,
                location="GLO",
                unit=unit,
                type="production",
            ),
            *exchanges,
        ],
    }


def electrolyte_inventories(source):
    s12 = source["polymer_electrolyte_S12"]
    peo = activity(
        "polyethylene oxide production, proxy, for ASSB",
        "polyethylene oxide, for ASSB",
        [technosphere("market for ethylene oxide", "ethylene oxide", 1, "RoW")],
        "Material proxy: one kg ethylene oxide per kg addition polymer. Corrects "
        "the polyethylene label in Zhang 2025 S12 to the named PEO chemistry. "
        "Polymerization energy and catalyst inventory remain a coverage gap; "
        "electrolyte mixing/heating/drying is supplied separately. "
        "https://doi.org/10.1016/j.jclepro.2025.146607",
    )
    polymer = activity(
        "polymer electrolyte production, PEO-LiTFSI, for ASSB",
        "polymer electrolyte, PEO-LiTFSI, for ASSB",
        [
            technosphere(
                peo["name"], peo["reference product"], s12["polyethylene_oxide_kg"]
            ),
            technosphere(
                "LiTFSi, (CF3SO2)2NLi, production, for LiS battery",
                "LiTFSi, (CF3SO2)2NLi",
                s12["LiTFSI_kg"],
                comment="Existing inventory for the same LiTFSI salt; manufacturing proxy.",
            ),
            technosphere(
                "market for argon, liquid", "argon, liquid", s12["argon_kg"], "RER"
            ),
            technosphere(
                "market for electricity, medium voltage",
                "electricity, medium voltage",
                s12["electricity_kWh"],
                "DE",
                "kilowatt hour",
                "S12: 0.002 + 0.083 + 0.019 + 1.687 kWh/kg electrolyte.",
            ),
        ],
        "Zhang 2025 S12, 80% PEO and 20% LiTFSI retained solids. Argon is process "
        "gas, excluded from retained mass. German electricity is an explicit "
        "supply-location choice. https://doi.org/10.1016/j.jclepro.2025.146607",
    )
    s13 = source["llzo_S13"]
    llzo_exchanges = [
        technosphere(
            f"market for {name}",
            name,
            s13[key],
            "RoW" if name.startswith("aluminium oxide") else "GLO",
        )
        for name, key in [
            ("lithium carbonate", "lithium_carbonate_kg"),
            ("lanthanum oxide", "lanthanum_oxide_kg"),
            ("zirconium oxide", "zirconium_oxide_kg"),
            ("aluminium oxide, non-metallurgical", "aluminium_oxide_kg"),
        ]
    ]
    llzo_exchanges += [
        technosphere(
            "market for electricity, medium voltage",
            "electricity, medium voltage",
            s13["electricity_kWh"],
            "DE",
            "kilowatt hour",
        ),
        technosphere(
            "market for transport, freight train",
            "transport, freight train",
            s13["rail_tkm"],
            "Europe without Switzerland",
            "ton kilometer",
        ),
        technosphere(
            "market for transport, freight, lorry 16-32 metric ton, EURO6",
            "transport, freight, lorry 16-32 metric ton, EURO6",
            s13["lorry_tkm"],
            "RER",
            "ton kilometer",
        ),
        technosphere(
            "market for chemical factory, organics",
            "chemical factory, organics",
            s13["chemical_factory_unit"],
            unit="unit",
        ),
        technosphere(
            "market for hazardous waste, for incineration",
            "hazardous waste, for incineration",
            -s13["waste_kg"],
            "RoW",
            comment="Unspecified S13 process waste: hazardous-incineration treatment proxy.",
        ),
        biosphere("Carbon dioxide, fossil", s13["carbon_dioxide_kg"], "air"),
        biosphere("Lithium I", s13["lithium_air_kg"], "air::urban air close to ground"),
    ]
    llzo = activity(
        "oxide electrolyte production, LLZO, for ASSB",
        "oxide electrolyte, LLZO, for ASSB",
        llzo_exchanges,
        "Zhang 2025 S13 synthesis inventory per kg Li7La3Zr2O12. Reagent masses "
        "include process inputs/losses; they are not retained elemental composition. "
        "Unspecified waste uses a documented hazardous-incineration proxy; "
        "electricity uses DE. https://doi.org/10.1016/j.jclepro.2025.146607",
    )
    return [peo, polymer, llzo]


def zhang_cell(chemistry, design, electrolyte):
    c = design["composition_g_per_cell"]
    mass = c["Total_mass"]
    density = c["Cell_capacity"] / mass  # Wh/g = kWh/kg
    feeds = [
        (
            (
                "market for lithium iron phosphate"
                if chemistry == "polymer"
                else "market for NMC811 oxide"
            ),
            "lithium iron phosphate" if chemistry == "polymer" else "NMC811 oxide",
            c["Cathode active material"],
            "RoW",
        ),
        (
            "lithium foil production, for LiS battery",
            "lithium foil",
            c["Anode active material"],
            "GLO",
        ),
        (
            "market for carbon black",
            "carbon black",
            c["Cathode conductive additive"],
            "GLO",
        ),
        (
            "market for polyvinylfluoride",
            "polyvinylfluoride",
            c["Cathode Binder"],
            "GLO",
        ),
        (
            "market for aluminium collector foil, for Li-ion battery",
            "aluminium collector foil, for Li-ion battery",
            c["Cathode current collector"],
            "GLO",
        ),
        (
            electrolyte["name"],
            electrolyte["reference product"],
            c["Electrolyte"],
            "GLO",
        ),
        (
            "market for aluminium, wrought alloy",
            "aluminium, wrought alloy",
            c["Casing_al_layer"],
            "GLO",
        ),
        (
            "market for polyethylene terephthalate, granulate, amorphous",
            "polyethylene terephthalate, granulate, amorphous",
            c["Casing_al_pet"],
            "GLO",
        ),
        (
            "market for polypropylene, granulate",
            "polypropylene, granulate",
            c["Casing_al_pp"],
            "GLO",
        ),
    ]
    exchanges = [technosphere(n, p, amount / mass, loc) for n, p, amount, loc in feeds]
    for process in ("sheet rolling, aluminium", "extrusion, plastic film"):
        material = (
            c["Casing_al_layer"]
            if process.startswith("sheet")
            else c["Casing_al_pet"] + c["Casing_al_pp"]
        )
        exchanges.append(
            technosphere(f"market for {process}", process, material / mass)
        )
    exchanges += [
        technosphere(
            "market for electricity, medium voltage",
            "electricity, medium voltage",
            design["manufacturing_electricity_kWh_per_kWh"] * density,
            "DE",
            "kilowatt hour",
        ),
        technosphere(
            "market group for heat, district or industrial, natural gas",
            "heat, district or industrial, natural gas",
            design["manufacturing_heat_kWh_per_kWh"] * density * 3.6,
            "RER",
            "megajoule",
        ),
    ]
    return activity(
        f"battery cell production, ASSB, {chemistry}",
        f"battery cell, ASSB, {chemistry}",
        exchanges,
        "Zhang 2025 baseline bipolar pouch reference design reconstructed from "
        "the pinned author model; density held constant at every year. Dry cell "
        "materials sum to 1 kg. Cathode active material * specific capacity * "
        "mean voltage defines capacity; no liquid electrolyte, graphite anode "
        "or copper anode collector. Existing lithium-foil and PVF-for-PVDF "
        "manufacturing proxies; unreported manufacturing scrap/solvent supply "
        "is not inferred. Pack hardware supplied separately. "
        "https://doi.org/10.1016/j.jclepro.2025.146607",
        **{"battery cell energy density": density},
    )


def sulfide_cell(source):
    cell = source["rietdorf_cell"]
    rows = source["rietdorf_scenario_1_rows"]
    cam = next(row["amount"] for row in rows if row["source_row"] == 44)
    capacity = (
        cam
        * cell["cathode_yield"]
        * cell["assembly_yield"]
        * cell["cell_yield"]
        * cell["cathode_capacity_Ah_per_kg"]
        * cell["voltage_V"]
        / 1000
    )
    normalization = 1 / capacity
    density = cell["cell_energy_Wh"] / 1000 / cell["total_mass_kg"]
    factor = normalization * density
    exchanges = []
    for row in rows:
        if row["step"] == "Nickel plating":
            continue  # Reconstruct from original area-based inventory below.
        amount = row["amount"] * factor
        comment = f"Rietdorf 2024 SM2 Scenario 1 row {row['source_row']}; per-source-kWh amount {row['amount']}; all flows normalized by {normalization:.16g} before converting to kg cell."
        if row["supplier"].startswith("resources"):
            exchanges.append(
                biosphere("Germanium", amount, "natural resource::in ground")
            )
            exchanges[-1]["comment"] = (
                comment
                + " Published germanium resource/heat proxy retained; no refined-metal supplier is inferred."
            )
            continue
        name, loc = row["supplier"].rsplit(" (", 1)
        loc = loc.removesuffix(")")
        # Explicit 3.10 supplier identities/geographies for the source's 3.9 names.
        if name == "electrolysis of lithium chloride":
            name = "lithium production, lithium chloride electrolysis"
        if (
            name == "market for sodium hydroxide, without water, in 50% solution state"
            and loc == "GLO"
        ):
            loc = "RoW"
        if name == "market for xylene":
            name = "market for xylene, mixed"
        if name.startswith(
            "polyethylene, high density, granulate, recycled to generic market"
        ):
            name, loc = "market for polyethylene, high density, granulate", "GLO"
            comment += (
                " Native HDPE market replaces the source's cutoff-specific recycled supplier. "
                "The original quantity is retained; the supplied market mix determines recycling content. "
                "The 3.10 global market is migrated to the target release's regional markets."
            )
        unit = {
            "kg": "kilogram",
            "kWh": "kilowatt hour",
            "MJ": "megajoule",
            "m3": "cubic meter",
        }[row["unit"]]
        # Native 3.10 NMC811 markets are CN/RoW, as named in the source.
        product = row["flow"]
        if product == "xylene":
            product = "xylene, mixed"
        if product == "Electricity":
            product = "electricity, medium voltage"
        if row["direction"].strip() == "Output":
            amount = -amount
        exchanges.append(technosphere(name, product, amount, loc, unit, comment))

    plating = source["plating_original_source"]
    area = plating["area_m2_per_source_kWh"] * factor
    nickel = (
        plating["thickness_m"]
        * plating["nickel_density_kg_m3"]
        / plating["bath_efficiency"]
    )
    for n, p, per_area, loc, unit in [
        ("market for nickel, class 1", "nickel, class 1", nickel, "GLO", "kilogram"),
        (
            "market for water, deionised",
            "water, deionised",
            plating["water_kg_m2"],
            "Europe without Switzerland",
            "kilogram",
        ),
        (
            "market for electricity, medium voltage",
            "electricity, medium voltage",
            plating["electricity_kWh_m2"],
            "DE",
            "kilowatt hour",
        ),
        (
            "market for wastewater, average",
            "wastewater, average",
            -plating["wastewater_m3_m2"],
            "Europe without Switzerland",
            "cubic meter",
        ),
        (
            "market for hazardous waste, for incineration",
            "hazardous waste, for incineration",
            -plating["sludge_kg_m2"],
            "Europe without Switzerland",
            "kilogram",
        ),
    ]:
        exchanges.append(
            technosphere(
                n,
                p,
                area * per_area,
                loc,
                unit,
                "Chordia 2021 S5-10/S5-11 per m2, multiplied by Rietdorf plated area. kg Ni/m2 = 1e-6 m * 8900 kg/m3 / 0.95. Wastewater converted from dm3 to m3.",
            )
        )
    for amount, categories in [
        (area * plating["nickel_air_kg_m2"], "air"),
        (area * plating["nickel_water_kg_m2"], "water"),
    ]:
        exchanges.append(biosphere("Nickel II", amount, categories))
        exchanges[-1][
            "comment"
        ] = "Chordia S5-11: mg converted to kg; nickel-in-sulfamate water emission uses the dissolved nickel-ion flow proxy."
    return activity(
        "battery cell production, ASSB, sulfidic",
        "battery cell, ASSB, sulfidic",
        exchanges,
        "Rietdorf 2024 SM1/SM2 Scenario 1 NMC811/LiGPS/LiPS/lithium reference "
        "cell; baseline fixed over time. All published gross input/output flows "
        "are normalized to exactly 1 kWh from cathode capacity and reported "
        "cathode/assembly/final-cell yields, then converted to kg cells. "
        "Cell retained mass and density follow SM1, not its inconsistent volumetric "
        "density. Source coating flows are replaced by original Chordia 2021 "
        "area-based process. LiGPS germanium resource/heat and acrylonitrile "
        "binder proxies are retained and unvalidated process coverage. "
        "Pack hardware supplied separately. https://doi.org/10.1007/s11367-024-02355-1",
        **{
            "battery cell energy density": density,
            "source capacity normalization": normalization,
        },
    )


def build_inventories():
    source = json.loads(SOURCE.read_text())
    inventories = electrolyte_inventories(source)
    cells = [
        zhang_cell("polymer", source["zhang_baseline"]["polymer"], inventories[1]),
        zhang_cell("oxidic", source["zhang_baseline"]["oxidic"], inventories[2]),
        sulfide_cell(source),
    ]
    inventories.extend(cells)
    original = ExcelImporter(INVENTORIES / "lci-batteries-NMC955-LTO.xlsx")
    original.data = [
        ds for ds in original.data if ds["name"] == "battery production, Li-ion, NMC955"
    ]
    migrate_import_db(original, "3.8", "3.10")
    (template,) = original.data
    hardware = [
        e
        for e in template["exchanges"]
        if e["type"] == "technosphere" and "battery cell" not in e["name"]
    ]
    for cell, chemistry in zip(cells, CHEMISTRIES):
        density = cell["battery cell energy density"]
        pack = activity(
            f"battery production, ASSB, {chemistry}",
            f"battery, ASSB, {chemistry}",
            [
                technosphere(cell["name"], cell["reference product"], CELL_FRACTION),
                *deepcopy(hardware),
            ],
            "ASSB cell plus non-cell hardware and pack assembly from the existing "
            "NMC955 pack: 0.72591 kg cell/kg pack. Explicit generic hardware "
            "proxy, not a measured ASSB pack design; no second conventional cell. "
            "Pack materials, BMS, module packaging, cooling and assembly retained.",
            **{"battery cell mass share": CELL_FRACTION},
        )
        mass = 1 / (CELL_FRACTION * density)
        pack_market = activity(
            f"market for battery, ASSB, {chemistry}",
            pack["reference product"],
            [technosphere(pack["name"], pack["reference product"], 1)],
            "One kg complete manufactured ASSB pack per kg sold; no extra cell fraction.",
        )
        capacity = activity(
            f"market for battery capacity, ASSB, {chemistry}",
            "electricity storage capacity",
            [
                technosphere(
                    pack_market["name"], pack_market["reference product"], mass
                ),
                technosphere(
                    "market for used Li-ion battery",
                    "used Li-ion battery",
                    -mass,
                    "GLO",
                    comment="Equal-mass generic Li-ion end-of-life treatment proxy; no ASSB-specific recovery yield claimed.",
                ),
            ],
            f"One nominal installed kWh = {mass:.16g} kg pack * {CELL_FRACTION} "
            f"kg cells/kg pack * {density:.16g} kWh/kg cell. Published baseline "
            "cell design held constant at all scenario years. Manufacturing and "
            "end-of-life pack masses match. Cycle life/replacements belong to "
            "the consuming vehicle; no chemistry-specific lifetime is inferred.",
            unit="kilowatt hour",
            **{
                "battery cell mass share": CELL_FRACTION,
                "battery cell energy density": density,
            },
        )
        inventories.extend([pack, pack_market, capacity])
    return inventories


def write_workbook(data):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Solid-state batteries"
    sheet.append(["database", "premise solid-state reference batteries"])
    sheet.append([])
    columns = [
        "name",
        "amount",
        "location",
        "unit",
        "type",
        "reference product",
        "categories",
        "comment",
    ]
    for ds in data:
        sheet.append(["Activity", ds["name"]])
        for key, value in ds.items():
            if key not in {"exchanges", "name"}:
                sheet.append([key, value])
        sheet.append(["Exchanges"])
        sheet.append(columns)
        for exchange in ds["exchanges"]:
            sheet.append(
                [
                    (
                        (exchange.get("product") or exchange.get("reference product"))
                        if c == "reference product"
                        else exchange.get(c)
                    )
                    for c in columns
                ]
            )
        sheet.append([])
    sheet.column_dimensions["A"].width = 75
    sheet.column_dimensions["B"].width = 26
    sheet.column_dimensions["F"].width = 55
    sheet.column_dimensions["H"].width = 100
    workbook.save(INVENTORIES / "lci-batteries-ASSB.xlsx")


def update_capacity_placeholders():
    """Replace unused zero-share slots without shifting formulas or cached values."""
    path = INVENTORIES / "lci-battery-capacity.xlsx"
    replacements = dict(
        zip(
            [
                "market for battery capacity, Li-ion, LTO",
                "market for battery capacity, Li-ion, LiMn2O4",
                "market for battery capacity, Sodium-Nickel-Chloride, Na-NiCl",
            ],
            [f"market for battery capacity, ASSB, {c}" for c in CHEMISTRIES],
        )
    )
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    ET.register_namespace("", ns["s"])
    with ZipFile(path) as archive:
        entries = [(info, archive.read(info.filename)) for info in archive.infolist()]
        shared = []
        if "xl/sharedStrings.xml" in archive.namelist():
            shared = [
                "".join(el.itertext())
                for el in ET.fromstring(archive.read("xl/sharedStrings.xml"))
            ]
    updated = []
    counts = dict.fromkeys(replacements.values(), 0)
    for info, content in entries:
        if info.filename == "xl/worksheets/sheet1.xml":
            root = ET.fromstring(content)
            for row in root.findall("s:sheetData/s:row", ns):
                cells = {c.attrib["r"].rstrip("0123456789"): c for c in row}
                if "A" not in cells:
                    continue
                a = cells["A"]
                if a.get("t") == "s":
                    value = shared[int(a.find("s:v", ns).text)]
                else:
                    value = "".join(t.text or "" for t in a.findall("s:is/s:t", ns))
                if value in replacements or value in replacements.values():
                    # These slots are forecast-inactive in the original workbook.
                    if float(cells["B"].find("s:v", ns).text) != 0:
                        raise ValueError(
                            f"Expected zero-share placeholder at {a.attrib['r']}"
                        )
                    replacement = replacements.get(value, value)
                    for child in list(a):
                        a.remove(child)
                    a.set("t", "inlineStr")
                    ET.SubElement(
                        ET.SubElement(a, f"{{{ns['s']}}}is"), f"{{{ns['s']}}}t"
                    ).text = replacement
                    counts[replacement] += 1
            content = ET.tostring(root, encoding="utf-8")
        updated.append((info, content))
    if set(counts.values()) != {4}:
        raise ValueError(
            f"Expected three placeholders in each of four mobile markets: {counts}"
        )
    with ZipFile(path, "w") as archive:
        for info, content in updated:
            archive.writestr(info, content)


if __name__ == "__main__":
    write_workbook(build_inventories())
    update_capacity_placeholders()
