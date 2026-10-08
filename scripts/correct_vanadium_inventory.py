"""Apply the documented 2.5.4 correction to an unmodified 2.5.3 workbook."""

import argparse
import hashlib
import json
import math
from pathlib import Path
import xml.etree.ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator

ROOT = Path(__file__).resolve().parents[1]
parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument(
    "source", type=Path, help="Unmodified vanadium workbook from tag v.2.5.3"
)
parser.add_argument(
    "--output",
    type=Path,
    default=ROOT / "premise/data/additional_inventories/lci-batteries-vanadium.xlsx",
)
parser.add_argument(
    "--report", type=Path, default=ROOT / "benchmarks/vanadium-inventory-exchanges.json"
)
args = parser.parse_args()
SOURCE = args.source
expected = "6ce897e076fd5bcd35c8aaf5e9097addb61f9bb6cf03177dab74c7b793453903"
if hashlib.sha256(SOURCE.read_bytes()).hexdigest() != expected:
    parser.error(
        "The input must be the unchanged workbook from v.2.5.3; refusing to allocate twice."
    )
DEST = args.output
MINE_SHARE = (1.53 * 0.075168516669567) / (1.53 * 0.075168516669567 + 0.46 * 0.176011)
# Both refinery prices refer to 2010 US dollars; one short ton is 2000 lb.
REFINERY_SHARE = 6.46 / (6.46 + 140 / 2000)
weber = "Weber et al. (2018), doi:10.1021/acs.est.8b02073, supporting information"
wb = load_workbook(SOURCE)
cached = load_workbook(SOURCE, data_only=True)
sheet = wb["Sheet1"]
values = cached["Sheet1"]
starts = [
    i for i in range(1, sheet.max_row + 1) if sheet.cell(i, 1).value == "Activity"
]
removed_rows = []
changes = []

for start, end in zip(starts, starts[1:] + [sheet.max_row + 2]):
    metadata = {
        sheet.cell(i, 1).value: (i, sheet.cell(i, 2).value)
        for i in range(start, end)
        if sheet.cell(i, 1).value
        in ("Activity", "location", "reference product", "comment", "production amount")
    }
    name = metadata["Activity"][1]
    location = metadata["location"][1]
    header = next(i for i in range(start, end) if sheet.cell(i, 1).value == "name")
    cols = {c.value: c.column for c in sheet[header] if c.value is not None}
    rows = [i for i in range(header + 1, end) if values.cell(i, cols["name"]).value]
    if location != "ZA":
        continue
    comment_col = cols.get("comment")
    if comment_col is None:
        comment_col = max(cols.values()) + 1
        sheet.cell(header, comment_col, "comment")
    production = next(
        i for i in rows if values.cell(i, cols["type"]).value == "production"
    )
    old_output = values.cell(production, cols["amount"]).value
    factor = 1 / old_output
    note = f"Premise 2.5.4: normalized to 1 kg reference product from an output of {old_output:g} kg; all non-production exchanges divided by that output. "
    if name == "vanadium bearing magnetite production":
        factor *= MINE_SHARE
        note = (
            f"Premise 2.5.4: attributional mining inventory, {weber}, Table S9. "
            "Joint outputs: 1.53 kg vanadium-bearing magnetite and 0.46 kg ilmenite. "
            f"Economic share of shared burdens: {MINE_SHARE:.16g}; "
            "formula = (1.53 * 0.075168516669567) / (1.53 * 0.075168516669567 + 0.46 * 0.176011). "
            "Prices are EUR2005/kg from the magnetite and ilmenite production exchanges of "
            "ecoinvent 3.12 cut-off, ilmenite - magnetite mine operation, GLO. "
            "Ordinary magnetite is a price proxy: no vanadium premium is assumed. "
            "The ilmenite substitution credit is removed. Iron and vanadium extraction are assigned "
            "to the magnetite output; titanium extraction is assigned to the separated ilmenite. "
            "These product-specific resource flows are not economically scaled. "
            "This is a fixed foreground assumption for all supported background versions. "
        )
    elif name == "vanadium pentoxide production":
        factor *= REFINERY_SHARE
        note += (
            f"Attributional co-product allocation, {weber}, Table S14: "
            "1 kg V2O5 and 1 kg sodium sulfate output; 0.5 kg sodium sulfate remains an input. "
            "The sodium-sulfate substitution credit is removed. "
            f"V2O5 economic share = 6.46 / (6.46 + 140 / 2000) = {REFINERY_SHARE:.16g}. "
            "USGS prices, both 2010 USD: V2O5 6.46 USD/lb (Mineral Commodity Summaries 2014, vanadium); "
            "sodium sulfate 140 USD/short ton (Mineral Commodity Summaries 2011, sodium sulfate). "
            "US prices are geographic proxies for the South African refinery. "
            "Iron scrap is a recyclable output under cut-off, with no substitution benefit; "
            "its existing recycling exchange and waste-treatment directions are retained. "
        )
    elif name == "vanadium slag production":
        note += (
            f"{weber}, Tables S12-S14: the inherited 0.1226 output encodes "
            "0.0613 / 0.5, i.e. 50% economic allocation to the vanadium stream. "
            "This allocation is already included and is not applied again. "
            "The legacy reference product is measured as kg V2O5 contained in slag, "
            "not kg bulk slag at 25% V2O5. Published upstream yields are retained; "
            "the source tables do not close the physical vanadium balance. "
        )
    else:
        note += "Published process yields are retained. "
    note += (
        "Lognormal uncertainties of changed exchanges are centered on the corrected absolute amount "
        "(loc = ln(abs(amount))); relative scale is retained. "
    )
    old_comment = metadata.get("comment")
    if old_comment:
        sheet.cell(
            old_comment[0], 2, note + "\n\nOriginal source notes: " + old_comment[1]
        )
    if "production amount" in metadata:
        sheet.cell(metadata["production amount"][0], 2, 1)
    for row in rows:
        excname = values.cell(row, cols["name"]).value
        kind = values.cell(row, cols["type"]).value
        old = values.cell(row, cols["amount"]).value
        categories = values.cell(row, cols.get("categories", 19)).value
        removed = (
            name == "vanadium bearing magnetite production"
            and (
                excname == "ilmenite - magnetite mine operation"
                or excname == "Titanium"
            )
        ) or (
            name == "vanadium pentoxide production"
            and excname == "market for sodium sulfate, anhydrite"
            and old < 0
        )
        if removed:
            removed_rows.append(row)
            changes.append(
                {
                    "activity": name,
                    "exchange": excname,
                    "before": old,
                    "after": None,
                    "reason": "co-product handled by allocation, without an avoided-production input",
                }
            )
            continue
        new = 1.0 if kind == "production" else old * factor
        reason = f"Original amount {old:.16g}; scale {factor:.16g}."
        if name == "vanadium bearing magnetite production" and excname in (
            "Iron",
            "Vanadium",
        ):
            new = old
            reason = "Element-specific extraction assigned to the magnetite output; no economic allocation of this flow."
        sheet.cell(row, cols["amount"], new)
        old_exc_comment = values.cell(row, comment_col).value or ""
        sheet.cell(
            row,
            comment_col,
            (old_exc_comment + "\n" if old_exc_comment else "")
            + "Premise 2.5.4: "
            + reason,
        )
        if "uncertainty type" in cols:
            uncertainty = values.cell(row, cols["uncertainty type"]).value
            if kind == "production":
                sheet.cell(row, cols["uncertainty type"], 0)
                if "loc" in cols:
                    sheet.cell(row, cols["loc"], 1.0)
            elif uncertainty == 2:
                sheet.cell(row, cols["loc"], math.log(abs(new)))
                if new < 0 and "negative" not in cols:
                    cols["negative"] = max(comment_col, max(cols.values())) + 1
                    sheet.cell(header, cols["negative"], "negative")
                if "negative" in cols:
                    sheet.cell(row, cols["negative"], new < 0)
        changes.append(
            {
                "activity": name,
                "exchange": excname,
                "before": old,
                "after": new,
                "reason": reason,
            }
        )

# Preserve formula caches, including the unmodified Chinese route and skipped notes.
# openpyxl writes formula cells without a cached result. Cache each original result,
# and translate retained formulas after row deletion (mostly LN(B<row>) uncertainties).
formula_caches = {}
for sh in wb:
    for row in sh:
        for cell in row:
            if cell.data_type == "f":
                value = cached[sh.title][cell.coordinate].value
                if value is None:
                    raise ValueError(
                        f"Missing original formula cache: {sh.title}!{cell.coordinate}"
                    )
                new_row = cell.row - (
                    sum(r < cell.row for r in removed_rows)
                    if sh.title == "Sheet1"
                    else 0
                )
                new_coordinate = f"{cell.column_letter}{new_row}"
                formula_caches[(sh.title, new_coordinate)] = value
                if new_row != cell.row:
                    cell.value = Translator(
                        cell.value, origin=cell.coordinate
                    ).translate_formula(new_coordinate)
for row in sorted(removed_rows, reverse=True):
    sheet.delete_rows(row)

assumptions = wb.create_sheet("Allocation 2.5.4")
for row in [
    ["skip"],
    [
        "Method",
        "Attributional allocation; product-specific metal resources assigned by element",
    ],
    ["Mining share, magnetite", MINE_SHARE],
    ["Mining share, ilmenite", 1 - MINE_SHARE],
    ["Magnetite price EUR2005/kg", 0.075168516669567],
    ["Ilmenite price EUR2005/kg", 0.176011],
    [
        "Mining price source",
        "ecoinvent 3.12 cut-off, ilmenite - magnetite mine operation, GLO, production-exchange price properties",
    ],
    [
        "Mining proxy limitation",
        "Ordinary magnetite price used for V-bearing magnetite; no vanadium premium. Original 77/23 sheet is historical and is not used.",
    ],
    ["Refinery share, V2O5", REFINERY_SHARE],
    ["Refinery share, sodium sulfate", 1 - REFINERY_SHARE],
    ["V2O5 price USD2010/lb", 6.46],
    ["Sodium sulfate price USD2010/short ton", 140],
    [
        "V2O5 price source",
        "https://d9-wret.s3.us-west-2.amazonaws.com/assets/palladium/production/mineral-pubs/vanadium/mcs-2014-vanad.pdf",
    ],
    [
        "Sodium sulfate price source",
        "https://d9-wret.s3.us-west-2.amazonaws.com/assets/palladium/production/mineral-pubs/sodium-sulfate/mcs-2011-nasul.pdf",
    ],
    [
        "Steel/slag allocation",
        "Original 50% retained; normalize 0.1226 kg accounting output to 1 kg V2O5 contained in slag.",
    ],
    [
        "Limitation",
        "Published ore-to-slag vanadium yields are inconsistent; this correction resolves allocation/substitution and normalization, not the source metallurgical balance.",
    ],
]:
    assumptions.append(row)
assumptions.column_dimensions["A"].width = 40
assumptions.column_dimensions["B"].width = 110
wb.save(DEST)

ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
vtag = "{" + ns["m"] + "}v"
with ZipFile(DEST) as archive:
    contents = {n: archive.read(n) for n in archive.namelist()}
for index, name in enumerate(wb.sheetnames, 1):
    path = f"xl/worksheets/sheet{index}.xml"
    root = ET.fromstring(contents[path])
    for cell in root.findall(".//m:c", ns):
        key = (name, cell.attrib["r"])
        if key in formula_caches and cell.find("m:f", ns) is not None:
            element = cell.find("m:v", ns)
            if element is None:
                element = ET.SubElement(cell, vtag)
            value = formula_caches[key]
            cell.set("t", "str" if isinstance(value, str) else "n")
            element.text = str(value)
    contents[path] = ET.tostring(root, encoding="utf-8")
with ZipFile(DEST, "w", ZIP_DEFLATED) as archive:
    for name, content in contents.items():
        archive.writestr(name, content)
args.report.write_text(
    json.dumps(
        {
            "mining_share": MINE_SHARE,
            "refinery_share": REFINERY_SHARE,
            "changes": changes,
        },
        indent=2,
    )
    + "\n"
)
print(
    json.dumps(
        {
            "mining_share": MINE_SHARE,
            "refinery_share": REFINERY_SHARE,
            "changed_exchanges": len(changes),
            "removed_rows": removed_rows,
        }
    )
)
