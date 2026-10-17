"""Reconcile VRFB electrolyte, nominal capacity and lifetime material quantities.

Offline authoring only. Preserve all unrelated Excel formula/cache payloads.
Sources and explicit reference-model assumptions are packaged in JSON.
"""

import json
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "premise/data"
INVENTORIES = DATA / "additional_inventories"
SOURCE = DATA / "battery/vanadium_reference_design.json"
CONDITIONING = "vanadium electrolyte conditioning, reference design"
USED_MARKET = "market for used redox-flow battery, Vanadium, reference design"
USED_PRODUCT = "used redox-flow battery, Vanadium, reference design"
NOTE = (
    "VRFB reference correction v1 (2026-10-17). Independent reconstruction: "
    "Oxkem VOS16 2024 electrolyte; Martin 2020 electrochemical reduction; "
    "Weber 2018 hardware and lifetime. Nominal capacity uses paired one-electron "
    "V charge at Li 2011's 1.25 V standard potential, an explicit reference "
    "assumption, not a measured discharge average. Legacy '8.3 MWh' activity "
    "name is an identifier, not the reconstructed capacity. The source archive's "
    "unspecified remaining-error warning is unresolved; no archive processes "
    "are copied. See data/battery/vanadium_reference_design.json and stationary "
    "battery documentation for calculations and coverage gaps. Fixed design "
    "at all years, no imposed primary/secondary-share change."
)


def reference():
    source = json.loads(SOURCE.read_text())
    e, c, h, life = (
        source[k] for k in ("electrolyte", "constants", "hardware_kg", "lifetime")
    )
    moles_V = e["vanadium_mol_L"] / e["density_kg_L"]
    metal = moles_V * c["V_kg_mol"]
    oxide = moles_V / 2 * (2 * c["V_kg_mol"] + 5 * c["O_kg_mol"])
    acid = e["sulfate_mol_L"] * c["H2SO4_kg_mol"] / e["density_kg_L"]
    phosphate = e["phosphoric_acid_mol_L"] * c["H3PO4_kg_mol"] / e["density_kg_L"]
    oxygen = moles_V * e["electrons_per_V_reduction"] / 4 * 2 * c["O_kg_mol"]
    carrier_water = phosphate * (1 / e["phosphoric_acid_purchase_mass_fraction"] - 1)
    water = 1 + oxygen - oxide - acid - phosphate - carrier_water
    stack = sum(h["stack"].values())
    periphery = sum(h["periphery"].values())
    mass = stack + periphery + h["electrolyte"] + h["tank"]
    # Both reservoirs contain V, but the paired half-reactions transfer one
    # electron per V PAIR, not per V atom in the complete system.
    charge_Ah = h["electrolyte"] * moles_V / 2 * c["faraday_C_mol"] / 3600
    capacity = charge_Ah * c["reference_cell_voltage_V"] / 1000
    cycles = life["years"] * life["days_per_year"] * life["cycles_per_day"]
    replacements = life["years"] / life["stack_years"] - 1
    return dict(
        source=source,
        oxide=oxide,
        acid=acid,
        phosphate=phosphate,
        oxygen=oxygen,
        carrier_water=carrier_water,
        water=water,
        metal=metal,
        electricity=metal * e["electricity_kWh_per_kg_V"],
        stack=stack,
        periphery=periphery,
        mass=mass,
        capacity=capacity,
        density=capacity / mass,
        mass_per_kWh=mass / capacity,
        charge_Ah=charge_Ah,
        cycles=cycles,
        replacements=replacements,
        lifetime_kWh=capacity * life["usable_SOC_fraction"] * cycles,
    )


def exchange(name, product, amount, unit="kilogram", location="RER", **extra):
    return dict(
        name=name,
        **{"reference product": product},
        amount=amount,
        unit=unit,
        location=location,
        type="technosphere",
        **extra,
    )


def auxiliary_inventories():
    r = reference()
    conditioning = dict(
        name=CONDITIONING,
        **{"reference product": CONDITIONING},
        location="RER",
        unit="unit",
        **{"production amount": 1},
        type="process",
        comment=NOTE
        + " One unit conditions one kg final solution, not one kg material. Water and electricity only here; V2O5/acids are bought by final electrolyte production. Reduction electricity is per kg V. Electrode/membrane wear, peripheral energy and process-factory inventory remain gaps.",
    )
    conditioning["exchanges"] = [
        {**exchange(CONDITIONING, CONDITIONING, 1, unit="unit"), "type": "production"},
        exchange(
            "market for water, deionised",
            "water, deionised",
            r["water"],
            location="Europe without Switzerland",
        ),
        exchange(
            "market for electricity, medium voltage",
            "electricity, medium voltage",
            r["electricity"],
            unit="kilowatt hour",
            location="DE",
            comment="Martin 2020 Table 6 / Eq.5: 1.69 kWh/kg elemental V; DE electricity is a location proxy. Excludes peripheral devices.",
        ),
        dict(
            name="Oxygen",
            amount=r["oxygen"],
            unit="kilogram",
            type="biosphere",
            categories="air",
            comment="Stoichiometric OER accompanying V(V) to mean V(III/IV): 1.5 electrons/V, four electrons/O2.",
        ),
    ]
    waste = dict(
        name=USED_MARKET,
        **{"reference product": USED_PRODUCT},
        location="GLO",
        unit="kilogram",
        **{"production amount": -1},
        type="process",
        comment=NOTE
        + " Per kg nominal installed system; includes treatment of one replacement stack over 20 years. No substitution or recovery-credit changes.",
    )
    waste["exchanges"] = [
        {
            **exchange(USED_MARKET, USED_PRODUCT, -1, location="GLO"),
            "type": "production",
        }
    ]
    for part, mass in [
        ("stack", (1 + r["replacements"]) * r["stack"]),
        ("electrolyte tank", r["source"]["hardware_kg"]["tank"]),
        ("electrolyte solution", r["source"]["hardware_kg"]["electrolyte"]),
        ("periphericals", r["periphery"]),
    ]:
        waste["exchanges"].append(
            exchange(
                f"treatment of {part}, for VRFB system",
                f"used {part}, for VRFB system",
                -mass / r["mass"],
            )
        )
    return [conditioning, waste]


def write_auxiliary(path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "VRFB reference"
    sheet.append(["database", "premise VRFB reference"])
    sheet.append([])
    cols = [
        "name",
        "amount",
        "location",
        "unit",
        "categories",
        "type",
        "reference product",
        "comment",
    ]
    # Conditioning precedes the main workbook; the waste market lives with
    # the four treatment recipes it references in that main workbook.
    for ds in auxiliary_inventories()[:1]:
        sheet.append(["Activity", ds["name"]])
        for key, value in ds.items():
            if key not in {"name", "exchanges"}:
                sheet.append([key, value])
        sheet.append(["Exchanges"])
        sheet.append(cols)
        for exc in ds["exchanges"]:
            sheet.append([exc.get(c) for c in cols])
        sheet.append([])
    if path.exists():
        old = load_workbook(path).active

        # Do not invalidate caches for a last-digit Excel serialization change.
        def signature(ws):
            return [
                tuple(format(v, ".14g") if isinstance(v, float) else v for v in row)
                for row in ws.values
            ]

        if signature(old) == signature(sheet):
            return False
    workbook.save(path)
    return True


def patch_cells(path, updates):
    """Replace only explicitly selected cells; retain other formulas/cached XML."""
    ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    ET.register_namespace("", ns)
    with ZipFile(path) as archive:
        entries = [(info, archive.read(info.filename)) for info in archive.infolist()]
    changed = False
    for part, cells in updates.items():
        i = next(i for i, (info, _) in enumerate(entries) if info.filename == part)
        info, content = entries[i]
        root = ET.fromstring(content)
        rows = {int(row.get("r")): row for row in root.iter(f"{{{ns}}}row")}
        by_address = {cell.get("r"): cell for cell in root.iter(f"{{{ns}}}c")}
        for address, value in cells.items():
            cell = by_address.get(address)
            if cell is None:
                number = int("".join(c for c in address if c.isdigit()))
                if number not in rows:
                    rows[number] = ET.SubElement(
                        root.find(f"{{{ns}}}sheetData"),
                        f"{{{ns}}}row",
                        {"r": str(number)},
                    )
                row = rows[number]
                cell = ET.SubElement(row, f"{{{ns}}}c", {"r": address})
            for child in list(cell):
                cell.remove(child)
            if value is None:
                cell.attrib.pop("t", None)
            elif isinstance(value, (int, float)):
                cell.attrib.pop("t", None)
                ET.SubElement(cell, f"{{{ns}}}v").text = repr(value)
            else:
                cell.set("t", "inlineStr")
                ET.SubElement(ET.SubElement(cell, f"{{{ns}}}is"), f"{{{ns}}}t").text = (
                    value
                )
        dimension = root.find(f"{{{ns}}}dimension")
        if dimension is not None:
            last = dimension.get("ref").split(":")[-1]
            old_row = int("".join(c for c in last if c.isdigit()))
            if max(rows) > old_row:
                dimension.set("ref", f"A1:H{max(rows)}")
        new = ET.tostring(root, encoding="utf-8")
        if new != content:
            changed = True
        entries[i] = (info, new)
    # A cached calculation chain must not retain references to formulas that
    # were deliberately replaced by constants in the targeted cells.
    contents = {info.filename: content for info, content in entries}
    if "xl/calcChain.xml" in contents:
        rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
        targets = {
            rel.get("Id"): rel.get("Target").lstrip("/")
            for rel in ET.fromstring(contents["xl/_rels/workbook.xml.rels"])
        }
        sheets = {}
        for sheet in ET.fromstring(contents["xl/workbook.xml"]).iter(f"{{{ns}}}sheet"):
            target = targets[sheet.get(f"{{{rel_ns}}}id")]
            sheets[sheet.get("sheetId")] = (
                target if target.startswith("xl/") else f"xl/{target}"
            )
        i = next(
            i
            for i, (info, _) in enumerate(entries)
            if info.filename == "xl/calcChain.xml"
        )
        info, content = entries[i]
        chain = ET.fromstring(content)
        sheet_id = None
        removed = False
        for cell in list(chain):
            sheet_id = cell.get("i", sheet_id)
            if cell.get("r") in updates.get(sheets.get(sheet_id), {}):
                chain.remove(cell)
                removed = True
            elif sheet_id is not None:
                cell.set("i", sheet_id)
        if removed:
            entries[i] = (info, ET.tostring(chain, encoding="utf-8"))
            changed = True
    if changed:
        with ZipFile(path, "w") as archive:
            for info, content in entries:
                archive.writestr(info, content)
    return changed


def patch_vrfb(path):
    r = reference()
    # Check the historical layout before using targeted XML cell addresses.
    values = load_workbook(path, data_only=True).active
    for row, name in [
        (95, "vanadium electrolyte solution production"),
        (406, "vanadium-redox flow battery system assembly, 8.3 megawatt hour"),
        (
            427,
            "electricity supply, high voltage, from vanadium-redox flow battery system",
        ),
        (446, "market for battery, redox-flow, Vanadium"),
        (461, "market for used redox-flow battery"),
    ]:
        if values.cell(row, 2).value != name:
            raise ValueError(
                f"Unexpected VRFB layout at row {row}; review the source update."
            )
    if (
        abs(values["B106"].value - 0.286) > 1e-10
        and abs(values["B106"].value - r["oxide"]) > 1e-10
    ):
        raise ValueError("Unknown electrolyte V2O5 dose; refusing a partial overwrite.")
    for address, legacy, corrected in [
        ("B107", 0.196, r["acid"]),
        ("B108", 0.004899, r["phosphate"]),
        ("B109", 0.863, 1),
    ]:
        if min(abs(values[address].value - v) for v in (legacy, corrected)) > 1e-10:
            raise ValueError(
                f"Unknown electrolyte coefficient {address}; review source update."
            )
    updates = dict(
        B98=NOTE,
        B106=r["oxide"],
        B107=r["acid"],
        B108=r["phosphate"],
        A109=CONDITIONING,
        B109=1,
        C109="RER",
        D109="unit",
        G109=CONDITIONING,
        H109="Conditioning service per kg final electrolyte: make-up water, V reduction electricity and O2 output. Materials remain separate.",
        B408=NOTE,
        B429=NOTE,
        B448=NOTE,
        B438=1 / r["source"]["lifetime"]["round_trip_efficiency"],
        B439=1 / r["lifetime_kWh"],
        B440=r["replacements"] / r["lifetime_kWh"],
        B441=-(1 + r["replacements"]) * r["stack"] / r["lifetime_kWh"],
        B442=-r["source"]["hardware_kg"]["tank"] / r["lifetime_kWh"],
        B443=-r["source"]["hardware_kg"]["electrolyte"] / r["lifetime_kWh"],
        B444=-r["periphery"] / r["lifetime_kWh"],
        H439="One installed system per reference lifetime delivered kWh, at 95% usable SOC; see reference design.",
        B457=1 / r["mass"],
        B458=r["replacements"] / r["mass"],
        H457="319950 kg = sum of S1 component rows (printed total differs by 2 kg).",
        H458="One replacement stack over 20 years; reference installed mass is not doubled.",
    )
    for row, key in [
        (110, "lorry"),
        (111, "train"),
        (112, "sea"),
        (113, "barge"),
        (114, "lorry_ZA"),
    ]:
        updates[f"B{row}"] = (
            r["oxide"]
            * r["source"]["original_electrolyte_transport_tkm_per_batch"][key]
            / r["source"]["original_transport_V2O5_batch_kg"]
        )
        updates[f"H{row}"] = (
            "Historical S8 route-distance proxy, scaled with the corrected V2O5 shipment mass, per kg final solution."
        )
    # Freeze the generic legacy other-chemistry treatment proxy: its formulas
    # referred to the standalone VRFB service being corrected above.
    for row, value in [
        (472, -0.06279145361121181),
        (473, -0.10107892532536537),
        (474, -0.8499818720542087),
        (475, -0.01531498868566142),
    ]:
        updates[f"B{row}"] = value
    # Local providers must be available during sequential inventory imports.
    starts = [
        row[0].row
        for row in values
        if row[0].value == "Activity" and row[1].value == USED_MARKET
    ]
    if len(starts) > 1:
        raise ValueError("Ambiguous VRFB reference waste market.")
    start = starts[0] if starts else values.max_row + 2
    waste = auxiliary_inventories()[1]
    rows = [["Activity", waste["name"]]]
    rows.extend(
        [key, value] for key, value in waste.items() if key not in {"name", "exchanges"}
    )
    columns = [
        "name",
        "amount",
        "location",
        "unit",
        "categories",
        "type",
        "reference product",
        "comment",
    ]
    rows.extend([["Exchanges"], columns])
    rows.extend([exc.get(c) for c in columns] for exc in waste["exchanges"])
    for offset, row in enumerate(rows):
        for column, value in enumerate(row):
            updates[f"{chr(65+column)}{start+offset}"] = value
    return patch_cells(path, {"xl/worksheets/sheet1.xml": updates})


def patch_capacity(path):
    r = reference()
    sheet = load_workbook(path, data_only=True)["Battery - stationary"]
    if (
        sheet["B80"].value
        != "market for battery capacity, redox-flow, Vanadium, stationary"
    ):
        raise ValueError("Unexpected VRFB capacity layout; review source update.")
    if (
        min(abs(sheet["B90"].value - v) for v in (1 / 0.0194, r["mass_per_kWh"]))
        > 1e-10
    ):
        raise ValueError("Unknown VRFB capacity coefficient; review source update.")
    updates = dict(
        B81=NOTE
        + f" One nominal gross reference kWh requires {r['mass_per_kWh']:.12g} kg installed system. No 21% extra BoP factor; hardware is already complete. Original capacity bounds mixed incompatible designs; deterministic reference replaces them. Actual discharge performance and generic consumer throughput remain separate limitations.",
        B90=r["mass_per_kWh"],
        G90=0,
        H90=r["mass_per_kWh"],
        I90=None,
        J90=None,
        A91=USED_MARKET,
        F91=USED_PRODUCT,
        B91=-r["mass_per_kWh"],
        G91=0,
        H91=-r["mass_per_kWh"],
        I91=None,
        J91=None,
    )
    # The sheet position is checked, rather than inferred from an activity name.
    return patch_cells(path, {"xl/worksheets/sheet2.xml": updates})


def main():
    print(
        "VRFB reference:",
        {k: v for k, v in reference().items() if k not in ("source",)},
    )
    print(
        "Auxiliary workbook changed:",
        write_auxiliary(INVENTORIES / "lci-vrfb-reference-conditioning.xlsx"),
    )
    print(
        "VRFB workbook changed:",
        patch_vrfb(INVENTORIES / "lci-batteries-vanadium-redox-flow.xlsx"),
    )
    print(
        "Capacity workbook changed:",
        patch_capacity(INVENTORIES / "lci-battery-capacity.xlsx"),
    )


if __name__ == "__main__":
    main()
