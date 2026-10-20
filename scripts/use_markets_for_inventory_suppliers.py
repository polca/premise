"""Use native markets for the affected vehicle scrap and ASSB HDPE suppliers.

Patch only string cells in the XLSX archive, retaining formula caches, amounts,
uncertainty and other workbook metadata. Source generators use the same markets.
"""

from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils.cell import (
    get_column_letter,
    column_index_from_string,
    coordinate_from_string,
)

ROOT = Path(__file__).resolve().parents[1]
INVENTORIES = ROOT / "premise/data/additional_inventories"
SCRAP = {
    "iron scrap, unsorted, Recycled Content cut-off": (
        "market for iron scrap, unsorted",
        "GLO",
    ),
    "aluminium scrap, new, Recycled Content cut-off": (
        "market for aluminium scrap, new",
        "RoW",
    ),
    "aluminium scrap, post-consumer, Recycled Content cut-off": (
        "market for aluminium scrap, post-consumer",
        "GLO",
    ),
}
HDPE = "polyethylene, high density, granulate, recycled to generic market for high density PE granulate"
FILES = (
    "lci-electric-car-mechanical-powertrain.xlsx",
    "lci-two_wheelers.xlsx",
    "lci-batteries-ASSB.xlsx",
)


def updates_for(path):
    workbook = load_workbook(path, read_only=True, data_only=False)
    updates = {}
    changed = 0
    for index, sheet in enumerate(workbook, 1):
        headers = {}
        edits = {}
        for row in sheet:
            if (
                row[0].value == "comment"
                and isinstance(row[1].value, str)
                and "Uses existing cutoff scrap destinations" in row[1].value
            ):
                edits[row[1].coordinate] = row[1].value.replace(
                    "Uses existing cutoff scrap destinations",
                    "Uses native scrap markets",
                )
                continue
            if row[0].value == "name" and any(cell.value == "amount" for cell in row):
                headers = {
                    cell.value: cell.column - 1
                    for cell in row
                    if cell.value is not None
                }
                continue
            if not headers or row[0].value not in {*SCRAP, HDPE}:
                continue
            if row[headers["type"]].value != "technosphere":
                continue
            old = row[0].value
            if old == HDPE:
                name, location = (
                    "market for polyethylene, high density, granulate",
                    "GLO",
                )
                comment = (
                    "Native HDPE market replaces the source cutoff-specific recycled supplier; "
                    "quantity retained, recycling content follows the market mix. "
                    "The 3.10 global market is migrated to the target release regional markets."
                )
            else:
                name, location = SCRAP[old]
                comment = "Native scrap market replaces the cutoff-specific destination; signed scrap quantity retained."
                if location == "RoW":
                    comment += " RoW is the non-European destination for this global equipment reference."
            edits[row[0].coordinate] = name
            edits[row[headers["location"]].coordinate] = location
            if "comment" in headers:
                cell = row[headers["comment"]]
                address = f"{get_column_letter(headers['comment'] + 1)}{row[0].row}"
                edits[address] = (
                    (str(cell.value) + " ") if cell.value else ""
                ) + comment
            changed += 1
        if edits:
            updates[f"xl/worksheets/sheet{index}.xml"] = edits
    workbook.close()
    return updates, changed


def patch(path):
    updates, changed = updates_for(path)
    if not updates:
        return 0
    with ZipFile(path) as archive:
        entries = [(info, archive.read(info.filename)) for info in archive.infolist()]
    ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    for index, (info, data) in enumerate(entries):
        if info.filename not in updates:
            continue
        tree = ET.fromstring(data)
        cells = {cell.get("r"): cell for cell in tree.iter(f"{{{ns}}}c")}
        for address, value in updates[info.filename].items():
            cell = cells.get(address)
            if cell is None:
                row_number = "".join(c for c in address if c.isdigit())
                parent = next(
                    r for r in tree.iter(f"{{{ns}}}row") if r.get("r") == row_number
                )
                cell = ET.Element(f"{{{ns}}}c", {"r": address})
                column = column_index_from_string(coordinate_from_string(address)[0])
                position = next(
                    (
                        i
                        for i, child in enumerate(parent)
                        if child.tag == f"{{{ns}}}c"
                        and column_index_from_string(
                            coordinate_from_string(child.get("r"))[0]
                        )
                        > column
                    ),
                    len(parent),
                )
                parent.insert(position, cell)
            assert cell.find(f"{{{ns}}}f") is None, address
            for child in list(cell):
                cell.remove(child)
            cell.set("t", "inlineStr")
            ET.SubElement(ET.SubElement(cell, f"{{{ns}}}is"), f"{{{ns}}}t").text = value
        entries[index] = (info, ET.tostring(tree, encoding="utf-8"))
    with ZipFile(path, "w") as archive:
        for info, data in entries:
            archive.writestr(info, data)
    return changed


if __name__ == "__main__":
    for filename in FILES:
        print(filename, patch(INVENTORIES / filename), "supplier exchanges changed")
