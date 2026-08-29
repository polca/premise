"""Audit DAC water balances in the packaged carbon-capture workbook."""

import math
from pathlib import Path

import yaml
from openpyxl import load_workbook

REPOSITORY = Path(__file__).resolve().parents[1]
WORKBOOK = (
    REPOSITORY
    / "premise"
    / "data"
    / "additional_inventories"
    / "lci-carbon-capture.xlsx"
)
DAC_WATER_ASSUMPTIONS = REPOSITORY / "premise" / "data" / "cdr" / "dac_water.yaml"
SOLVENT_TEXT = "with a solvent-based direct air capture system"
SORBENT_TEXT = "with a sorbent-based direct air capture system"


def activities():
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    worksheet = workbook["DAC"]
    rows = list(worksheet.iter_rows(values_only=True))
    output = {}
    for index, row in enumerate(rows):
        if not row or row[0] != "Activity":
            continue
        end = next(
            (
                next_index
                for next_index in range(index + 1, len(rows))
                if rows[next_index] and rows[next_index][0] == "Activity"
            ),
            len(rows),
        )
        block = rows[index + 1 : end]
        header_index = next(
            block_index
            for block_index, block_row in enumerate(block)
            if block_row and block_row[0] == "name"
        )
        header = block[header_index]
        output[row[1]] = [
            {
                field: value
                for field, value in zip(header, exchange)
                if field is not None and value is not None
            }
            for exchange in block[header_index + 1 :]
            if exchange and exchange[0]
        ]
    workbook.close()
    return output


def one_exchange(exchanges, name, exchange_type=None):
    matches = [
        exchange
        for exchange in exchanges
        if exchange.get("name") == name
        and (exchange_type is None or exchange.get("type") == exchange_type)
    ]
    assert len(matches) == 1, (name, exchange_type, len(matches))
    return matches[0]


def main():
    inventory = activities()
    with open(DAC_WATER_ASSUMPTIONS, encoding="utf-8") as stream:
        assumptions = yaml.safe_load(stream)

    solvent = {
        name: exchanges
        for name, exchanges in inventory.items()
        if name.startswith("carbon dioxide, captured") and SOLVENT_TEXT in name
    }
    sorbent = {
        name: exchanges
        for name, exchanges in inventory.items()
        if name.startswith("carbon dioxide, captured") and SORBENT_TEXT in name
    }
    assert len(solvent) == 8
    assert len(sorbent) == 10

    expected = assumptions["solvent_based"]["operation"]
    for exchanges in solvent.values():
        make_up = one_exchange(exchanges, "market for tap water", "technosphere")
        evaporation = one_exchange(exchanges, "Water", "biosphere")
        assert math.isclose(make_up["amount"], expected["makeup_water_kg_per_kg_co2"])
        assert evaporation["categories"] == "air"
        assert math.isclose(make_up["amount"], 1000 * evaporation["amount"])

    for exchanges in sorbent.values():
        assert not any(
            exchange.get("name")
            in {"market for tap water", "market for water, deionised", "Water"}
            for exchange in exchanges
        )

    plant = inventory["direct air capture system, solvent-based, 1MtCO2"]
    initial_charge = one_exchange(plant, "market for tap water", "technosphere")
    assert (
        initial_charge["amount"]
        == assumptions["solvent_based"]["construction"]["initial_charge_kg_per_plant"]
    )

    sorbent_material = inventory[
        "amine-based silica production, for sorbent-based direct air capture system"
    ]
    silica = one_exchange(
        sorbent_material,
        "silica gel production, for sorbent-based direct air capture system",
    )
    resin_eol = one_exchange(
        sorbent_material,
        "treatment of spent anion exchange resin from potable water production, municipal incineration",
    )
    assert math.isclose(silica["amount"], 0.032)
    assert math.isclose(resin_eol["amount"], -0.36)

    pei = inventory[
        "polyethyleneimine (PEI) production, for sorbent-based direct air capture system"
    ]
    sodium_sulfate = one_exchange(pei, "market for sodium sulfate, anhydrite")
    assert math.isclose(sodium_sulfate["amount"], -4.595)

    silica_production = inventory[
        "silica gel production, for sorbent-based direct air capture system"
    ]
    silica_water = one_exchange(silica_production, "market for water, deionised")
    wastewater = one_exchange(
        silica_production, "treatment of wastewater, average, wastewater treatment"
    )
    assert math.isclose(silica_water["amount"], 40)
    assert math.isclose(wastewater["amount"], -0.035)

    print("DAC water inventory audit passed.")
    print("  solvent operating variants: 8; 3.4 kg in = 0.0034 m3 to air")
    print("  solvent initial charge: 744,000,000 kg/plant")
    print("  sorbent operating variants: 10; no external direct water")
    print("  virgin silica: 0.032 kg/kg sorbent; spent PEI proxy: 0.36 kg")
    print(
        "  explicitly unresolved: solvent EoL charge, PEI output water, 5 kg/kg silica"
    )


if __name__ == "__main__":
    main()
