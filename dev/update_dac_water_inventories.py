"""Apply the reviewed DAC water-accounting corrections to the source workbook.

This migration is intentionally idempotent. It edits the packaged Excel source only;
it does not import inventories, rebuild a Brightway database, or calculate LCIA.
"""

import argparse
import os
from copy import copy
from pathlib import Path
from tempfile import NamedTemporaryFile

from openpyxl import load_workbook

REPOSITORY = Path(__file__).resolve().parents[1]
WORKBOOK = (
    REPOSITORY
    / "premise"
    / "data"
    / "additional_inventories"
    / "lci-carbon-capture.xlsx"
)
SHEET = "DAC"
SOLVENT_OPERATION_TEXT = "with a solvent-based direct air capture system"
SORBENT_OPERATION_TEXT = "with a sorbent-based direct air capture system"
SOLVENT_PLANT = "direct air capture system, solvent-based, 1MtCO2"
SOLVENT_PLANT_EOL = "treatment of direct air capture system, solvent-based, 1MtCO2"
SORPTION_MATERIAL = (
    "amine-based silica production, for sorbent-based direct air capture system"
)
PEI_PRODUCTION = (
    "polyethyleneimine (PEI) production, for sorbent-based direct air capture system"
)
SILICA_PRODUCTION = "silica gel production, for sorbent-based direct air capture system"

SOLVENT_OPERATION_COMMENT = (
    " Water accounting: the Qiu et al. supplement reports 3.4 Mt/year cooling-"
    "tower make-up for a 1 MtCO2/year plant at 20 degrees C, 60% relative humidity "
    "and 90% operation. This is 3.4 kg make-up water/kg CO2 and is paired with "
    "0.0034 m3 Water-to-air/kg CO2 as operational evaporation. The separately "
    "reported 744,000-tonne initial charge is assigned to plant construction."
)
SORBENT_OPERATION_COMMENT = (
    " Water accounting: the Deutz and Bardow foreground reports no external "
    "operational water input. Atmospheric water co-adsorption/recovery is not "
    "credited in the base inventory because the quantity, site dependence and "
    "product-water quality are not specified."
)


def activity_blocks(worksheet):
    """Return ``(name, start, end)`` records for all activity blocks."""

    starts = [
        (row, worksheet.cell(row, 2).value)
        for row in range(1, worksheet.max_row + 1)
        if worksheet.cell(row, 1).value == "Activity"
    ]
    return [
        (
            name,
            row,
            starts[index + 1][0] - 1 if index + 1 < len(starts) else worksheet.max_row,
        )
        for index, (row, name) in enumerate(starts)
    ]


def block_for(worksheet, activity_name):
    matches = [
        block for block in activity_blocks(worksheet) if block[0] == activity_name
    ]
    if len(matches) != 1:
        raise ValueError(
            f"Expected one activity named {activity_name!r}; found {len(matches)}"
        )
    return matches[0]


def exchange_rows(worksheet, start, end):
    header = next(
        row for row in range(start, end + 1) if worksheet.cell(row, 1).value == "name"
    )
    return range(header + 1, end + 1)


def exchange_row(worksheet, activity_name, exchange_name):
    _, start, end = block_for(worksheet, activity_name)
    matches = [
        row
        for row in exchange_rows(worksheet, start, end)
        if worksheet.cell(row, 1).value == exchange_name
    ]
    if len(matches) != 1:
        raise ValueError(
            f"Expected one {exchange_name!r} exchange in {activity_name!r}; "
            f"found {len(matches)}"
        )
    return matches[0]


def metadata_row(worksheet, activity_name, field):
    _, start, end = block_for(worksheet, activity_name)
    return next(
        row for row in range(start, end + 1) if worksheet.cell(row, 1).value == field
    )


def append_once(value, addition):
    value = value or ""
    return value if addition.strip() in value else value.rstrip() + addition


def copy_row_style(worksheet, source, target):
    worksheet.row_dimensions[target].height = worksheet.row_dimensions[source].height
    for column in range(1, worksheet.max_column + 1):
        old = worksheet.cell(source, column)
        new = worksheet.cell(target, column)
        if old.has_style:
            new._style = copy(old._style)
        new.number_format = old.number_format
        new.protection = copy(old.protection)
        new.alignment = copy(old.alignment)


def set_exchange(worksheet, row, values):
    for column in range(1, worksheet.max_column + 1):
        worksheet.cell(row, column).value = None
    for column, value in enumerate(values, start=1):
        worksheet.cell(row, column).value = value


def ensure_exchange_after(worksheet, activity_name, after_name, exchange):
    """Insert an exchange after another one, or update the existing exchange."""

    _, start, end = block_for(worksheet, activity_name)
    existing = [
        row
        for row in exchange_rows(worksheet, start, end)
        if worksheet.cell(row, 1).value == exchange[0]
        and worksheet.cell(row, 6).value == exchange[5]
        and worksheet.cell(row, 5).value == exchange[4]
    ]
    if len(existing) > 1:
        raise ValueError(f"Duplicate {exchange[0]!r} exchanges in {activity_name!r}")
    if existing:
        set_exchange(worksheet, existing[0], exchange)
        return

    after = exchange_row(worksheet, activity_name, after_name)
    worksheet.insert_rows(after + 1)
    copy_row_style(worksheet, after, after + 1)
    set_exchange(worksheet, after + 1, exchange)


def update_solvent_operation(worksheet, activity_name):
    water = exchange_row(worksheet, activity_name, "market for tap water")
    worksheet.cell(water, 2).value = 3.4
    worksheet.cell(water, 8).value = (
        "DAC water: operational make-up replacing cooling-tower evaporation; "
        "3.4 Mt/year / 1 MtCO2/year = 3.4 kg/kg CO2 (Qiu et al. 2022 SI)."
    )
    ensure_exchange_after(
        worksheet,
        activity_name,
        "market for tap water",
        (
            "Water",
            0.0034,
            None,
            "cubic meter",
            "air",
            "biosphere",
            None,
            "DAC water: operational evaporation to air; 3.4 kg / "
            "1000 kg/m3 = 0.0034 m3/kg CO2.",
        ),
    )
    comment = metadata_row(worksheet, activity_name, "comment")
    worksheet.cell(comment, 2).value = append_once(
        worksheet.cell(comment, 2).value, SOLVENT_OPERATION_COMMENT
    )


def update_solvent_plant(worksheet):
    ensure_exchange_after(
        worksheet,
        SOLVENT_PLANT,
        "direct air capture system, solvent-based, 1MtCO2",
        (
            "market for tap water",
            744000000,
            "Europe without Switzerland",
            "kilogram",
            None,
            "technosphere",
            "tap water",
            "DAC water: initial circulating-water charge; 31,000 t/hour * 24 "
            "hours = 744,000 t/plant (Qiu et al. 2022 SI).",
        ),
    )
    comment = metadata_row(worksheet, SOLVENT_PLANT, "comment")
    worksheet.cell(comment, 2).value = append_once(
        worksheet.cell(comment, 2).value,
        " Water accounting: the 744,000-tonne initial circulating-water charge is "
        "modelled as plant capital stock. At 5e-11 plant/kg CO2, it contributes "
        "0.0372 kg water/kg CO2 over the 20-year, 1 Mt/year plant life. Qiu et al. "
        "do not specify the residual solution fate at decommissioning.",
    )
    eol_comment = metadata_row(worksheet, SOLVENT_PLANT_EOL, "comment")
    worksheet.cell(eol_comment, 2).value = append_once(
        worksheet.cell(eol_comment, 2).value,
        " Water accounting: the fate of the initial circulating solution at "
        "decommissioning is unresolved in the source. No unsupported water release "
        "or wastewater-treatment exchange is assumed.",
    )


def update_sorbent_operation(worksheet, activity_name):
    comment = metadata_row(worksheet, activity_name, "comment")
    worksheet.cell(comment, 2).value = append_once(
        worksheet.cell(comment, 2).value, SORBENT_OPERATION_COMMENT
    )


def update_sorbent_material(worksheet):
    silica = exchange_row(worksheet, SORPTION_MATERIAL, SILICA_PRODUCTION)
    worksheet.cell(silica, 2).value = 0.032
    worksheet.cell(silica, 8).value = (
        "Virgin silica requirement after 95% support recycling: "
        "0.64 kg * (1 - 0.95) = 0.032 kg/kg sorbent."
    )

    treatment = exchange_row(
        worksheet,
        SORPTION_MATERIAL,
        "treatment of spent anion exchange resin from potable water production, "
        "municipal incineration",
    )
    worksheet.cell(treatment, 2).value = -0.36
    worksheet.cell(treatment, 8).value = (
        "EoL proxy for the non-recycled 36% PEI fraction only; the 95% recycled "
        "silica support is represented by reduced virgin silica input."
    )

    comment = metadata_row(worksheet, SORPTION_MATERIAL, "comment")
    worksheet.cell(comment, 2).value = (
        "Amine-functionalized silica contains 0.36 kg polyethyleneimine (PEI) and "
        "0.64 kg silica support per kg sorbent. Following Qiu et al. (2022), 95% "
        "of the silica support is recycled, represented with the cut-off approach "
        "as 0.032 kg virgin silica input. The 0.36 kg PEI fraction is sent to the "
        "spent-resin incineration proxy. Solvent recovery is already represented "
        "inside the PEI-production inventory."
    )


def update_pei(worksheet):
    ensure_exchange_after(
        worksheet,
        PEI_PRODUCTION,
        "market for water, deionised",
        (
            "market for sodium sulfate, anhydrite",
            -4.595,
            "RoW",
            "kilogram",
            None,
            "technosphere",
            "sodium sulfate, anhydrite",
            "Co-product credit by system expansion. Midpoint of the Deutz and "
            "Bardow range: -(3.30 + 5.89) / 2 = -4.595 kg/kg PEI.",
            5,
            -4.595,
            -5.89,
            -3.3,
        ),
    )
    water = exchange_row(worksheet, PEI_PRODUCTION, "market for water, deionised")
    worksheet.cell(water, 8).value = (
        "Process-water input from Deutz and Bardow. The source does not report a "
        "complete output-water split; the water balance remains unresolved."
    )
    comment = metadata_row(worksheet, PEI_PRODUCTION, "comment")
    worksheet.cell(comment, 2).value = append_once(
        worksheet.cell(comment, 2).value,
        " Sodium sulfate is restored as a co-product credit (3.30-5.89 kg/kg PEI; "
        "midpoint 4.595 kg/kg). The reported deionised-water input is retained, "
        "while its output-water split is explicitly unresolved.",
    )


def update_silica(worksheet):
    water = exchange_row(worksheet, SILICA_PRODUCTION, "market for water, deionised")
    worksheet.cell(water, 8).value = (
        "40 kg process-water input per kg silica gel (Roes et al. inventory as "
        "reported by Deutz and Bardow)."
    )
    wastewater = exchange_row(
        worksheet,
        SILICA_PRODUCTION,
        "treatment of wastewater, average, wastewater treatment",
    )
    worksheet.cell(wastewater, 8).value = (
        "35 kg wastewater/kg silica converted at 1000 kg/m3 to 0.035 m3/kg "
        "silica. Qiu et al.'s 1.12 m3/kg sorbent is dimensionally inconsistent; "
        "35 * 0.64 * 0.05 = 1.12 kg = 0.00112 m3/kg sorbent."
    )
    comment = metadata_row(worksheet, SILICA_PRODUCTION, "comment")
    worksheet.cell(comment, 2).value = append_once(
        worksheet.cell(comment, 2).value,
        " Water accounting: 40 kg water enters and 35 kg leaves as wastewater per "
        "kg virgin silica. The source does not identify the remaining 5 kg; it is "
        "kept as an explicit unresolved balance rather than assigned to air or water.",
    )


def workbook_value_signature(workbook):
    """Return a deterministic signature of workbook structure and cell values."""

    return tuple(
        (
            worksheet.title,
            worksheet.max_row,
            worksheet.max_column,
            tuple(
                tuple(cell.value for cell in row)
                for row in worksheet.iter_rows(
                    min_row=1,
                    max_row=worksheet.max_row,
                    min_col=1,
                    max_col=worksheet.max_column,
                )
            ),
        )
        for worksheet in workbook.worksheets
    )


def migrate(workbook_path):
    workbook = load_workbook(workbook_path)
    signature_before = workbook_value_signature(workbook)
    worksheet = workbook[SHEET]

    blocks = activity_blocks(worksheet)
    solvent_activities = sorted(
        [
            (start, name)
            for name, start, _ in blocks
            if str(name).startswith("carbon dioxide, captured")
            and SOLVENT_OPERATION_TEXT in str(name)
        ],
        reverse=True,
    )
    sorbent_activities = [
        name
        for name, _, _ in blocks
        if str(name).startswith("carbon dioxide, captured")
        and SORBENT_OPERATION_TEXT in str(name)
    ]
    if len(solvent_activities) != 8:
        raise ValueError(
            f"Expected 8 solvent-DAC operating activities; found {len(solvent_activities)}"
        )
    if len(sorbent_activities) != 10:
        raise ValueError(
            f"Expected 10 sorbent-DAC operating activities; found {len(sorbent_activities)}"
        )

    # Process bottom-up so inserted rows cannot invalidate pending row positions.
    for _, activity_name in solvent_activities:
        update_solvent_operation(worksheet, activity_name)
    update_solvent_plant(worksheet)
    for activity_name in sorbent_activities:
        update_sorbent_operation(worksheet, activity_name)
    update_sorbent_material(worksheet)
    update_pei(worksheet)
    update_silica(worksheet)

    changed = signature_before != workbook_value_signature(workbook)
    if not changed:
        workbook.close()
        return len(solvent_activities), len(sorbent_activities), False

    with NamedTemporaryFile(
        suffix=workbook_path.suffix, dir=workbook_path.parent, delete=False
    ) as temporary:
        temporary_path = Path(temporary.name)
    try:
        workbook.save(temporary_path)
        os.replace(temporary_path, workbook_path)
    finally:
        if temporary_path.exists():
            temporary_path.unlink()
    workbook.close()
    return len(solvent_activities), len(sorbent_activities), True


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path, default=WORKBOOK)
    args = parser.parse_args()
    solvent_count, sorbent_count, changed = migrate(args.workbook)
    status = "Updated" if changed else "Already current"
    print(
        f"{status}: {args.workbook}: {solvent_count} solvent-DAC and "
        f"{sorbent_count} sorbent-DAC operating variants."
    )


if __name__ == "__main__":
    main()
