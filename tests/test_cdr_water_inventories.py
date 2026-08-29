"""Source-level tests for DAC water accounting; no database build is required."""

from copy import deepcopy

import pytest
import yaml
from openpyxl import load_workbook

from premise.carbon_dioxide_removal import (
    DAC_WATER_ASSUMPTIONS,
    CarbonDioxideRemoval,
)
from premise.filesystem_constants import INVENTORY_DIR

CDR_INVENTORY = INVENTORY_DIR / "lci-carbon-capture.xlsx"
SOLVENT_TEXT = "with a solvent-based direct air capture system"
SORBENT_TEXT = "with a sorbent-based direct air capture system"


@pytest.fixture(scope="module")
def inventory_activities():
    workbook = load_workbook(CDR_INVENTORY, read_only=True, data_only=True)
    worksheet = workbook["DAC"]
    rows = list(worksheet.iter_rows(values_only=True))
    activities = {}
    for index, row in enumerate(rows):
        if not row or row[0] != "Activity":
            continue
        end = next(
            (
                next_index
                for next_index in range(index + 1, len(rows))
                if rows[next_index] and rows[next_index][0] == "Activity"
            ),
            len(rows),
        )
        block = rows[index + 1 : end]
        metadata = {
            block_row[0]: block_row[1]
            for block_row in block
            if block_row and block_row[0] in {"comment", "source"}
        }
        header_index = next(
            block_index
            for block_index, block_row in enumerate(block)
            if block_row and block_row[0] == "name"
        )
        header = block[header_index]
        exchanges = [
            {
                field: value
                for field, value in zip(header, exchange)
                if field is not None and value is not None
            }
            for exchange in block[header_index + 1 :]
            if exchange and exchange[0]
        ]
        activities[row[1]] = {"metadata": metadata, "exchanges": exchanges}
    workbook.close()
    return activities


def one_exchange(activity, name, exchange_type=None):
    matches = [
        exchange
        for exchange in activity["exchanges"]
        if exchange.get("name") == name
        and (exchange_type is None or exchange.get("type") == exchange_type)
    ]
    assert len(matches) == 1
    return matches[0]


def test_water_assumptions_are_mass_consistent():
    with open(DAC_WATER_ASSUMPTIONS, encoding="utf-8") as stream:
        assumptions = yaml.safe_load(stream)

    operation = assumptions["solvent_based"]["operation"]
    construction = assumptions["solvent_based"]["construction"]
    assert operation["makeup_water_kg_per_kg_co2"] == pytest.approx(
        operation["evaporation_to_air_m3_per_kg_co2"]
        * assumptions["water_density_kg_per_m3"]
    )
    assert construction["annualized_initial_charge_kg_per_kg_co2"] == pytest.approx(
        construction["initial_charge_kg_per_plant"]
        * construction["plant_input_per_kg_co2"]
    )


def test_all_solvent_operating_variants_close_water_balance(inventory_activities):
    activities = [
        activity
        for name, activity in inventory_activities.items()
        if name.startswith("carbon dioxide, captured") and SOLVENT_TEXT in name
    ]
    assert len(activities) == 8
    for activity in activities:
        make_up = one_exchange(activity, "market for tap water", "technosphere")
        evaporation = one_exchange(activity, "Water", "biosphere")
        assert make_up["amount"] == pytest.approx(3.4)
        assert evaporation["amount"] == pytest.approx(0.0034)
        assert evaporation["categories"] == "air"
        assert make_up["amount"] == pytest.approx(1000 * evaporation["amount"])
        assert "744,000-tonne initial charge" in activity["metadata"]["comment"]


def test_solvent_initial_charge_is_capital_stock(inventory_activities):
    plant = inventory_activities["direct air capture system, solvent-based, 1MtCO2"]
    charge = one_exchange(plant, "market for tap water", "technosphere")
    assert charge["amount"] == 744_000_000
    assert "capital stock" in plant["metadata"]["comment"]

    treatment = inventory_activities[
        "treatment of direct air capture system, solvent-based, 1MtCO2"
    ]
    assert (
        "fate of the initial circulating solution" in treatment["metadata"]["comment"]
    )


def test_sorbent_operating_variants_have_no_external_direct_water(
    inventory_activities,
):
    activities = [
        activity
        for name, activity in inventory_activities.items()
        if name.startswith("carbon dioxide, captured") and SORBENT_TEXT in name
    ]
    assert len(activities) == 10
    for activity in activities:
        assert not any(
            exchange.get("name")
            in {"market for tap water", "market for water, deionised", "Water"}
            for exchange in activity["exchanges"]
        )
        assert "no external operational water input" in activity["metadata"]["comment"]


def test_sorbent_manufacture_applies_recycling_and_co_product_credit(
    inventory_activities,
):
    material = inventory_activities[
        "amine-based silica production, for sorbent-based direct air capture system"
    ]
    silica = one_exchange(
        material, "silica gel production, for sorbent-based direct air capture system"
    )
    treatment = one_exchange(
        material,
        "treatment of spent anion exchange resin from potable water production, "
        "municipal incineration",
    )
    assert silica["amount"] == pytest.approx(0.64 * (1 - 0.95))
    assert treatment["amount"] == pytest.approx(-0.36)

    pei = inventory_activities[
        "polyethyleneimine (PEI) production, for sorbent-based direct air capture system"
    ]
    sodium_sulfate = one_exchange(pei, "market for sodium sulfate, anhydrite")
    assert sodium_sulfate["amount"] == pytest.approx(-(3.30 + 5.89) / 2)
    assert sodium_sulfate["minimum"] == pytest.approx(-5.89)
    assert sodium_sulfate["maximum"] == pytest.approx(-3.30)


def test_silica_water_conversion_and_unresolved_balance_are_documented(
    inventory_activities,
):
    silica = inventory_activities[
        "silica gel production, for sorbent-based direct air capture system"
    ]
    water = one_exchange(silica, "market for water, deionised")
    wastewater = one_exchange(
        silica, "treatment of wastewater, average, wastewater treatment"
    )
    assert water["amount"] == pytest.approx(40)
    assert wastewater["amount"] == pytest.approx(-35 / 1000)
    assert "remaining 5 kg" in silica["metadata"]["comment"]
    assert "0.00112 m3/kg sorbent" in wastewater["comment"]


def solvent_dataset():
    return {
        "name": "carbon dioxide, captured, with a solvent-based direct air capture system",
        "reference product": "carbon dioxide, captured",
        "location": "EUR",
        "unit": "kilogram",
        "exchanges": [
            {
                "name": "market for tap water",
                "product": "tap water",
                "amount": 3.4,
                "unit": "kilogram",
                "type": "technosphere",
                "comment": "DAC water: operational make-up",
            },
            {
                "name": "Water",
                "amount": 0.0034,
                "unit": "cubic meter",
                "categories": ("air",),
                "type": "biosphere",
                "comment": "DAC water: operational evaporation",
            },
        ],
    }


def test_dac_water_scaling_is_paired_bounded_and_idempotent():
    cdr = object.__new__(CarbonDioxideRemoval)
    dataset = solvent_dataset()
    technology = "direct air capture (solvent, high-temp)"

    cdr.adjust_dac_water(dataset, technology, scaling_factor=1.2)
    first = deepcopy(dataset)
    cdr.adjust_dac_water(dataset, technology, scaling_factor=1.2)
    assert dataset == first
    assert dataset["exchanges"][0]["amount"] == pytest.approx(3.4 * 1.2)
    assert dataset["exchanges"][1]["amount"] == pytest.approx(0.0034 * 1.2)

    cdr.adjust_dac_water(dataset, technology, scaling_factor=0.1)
    assert dataset["exchanges"][0]["amount"] == pytest.approx(2.7)
    assert dataset["exchanges"][1]["amount"] == pytest.approx(0.0027)
    assert dataset["log parameters"]["DAC operational water scaling factor"] == (
        pytest.approx(2.7 / 3.4)
    )


def test_dac_water_scaling_rejects_an_unpaired_exchange():
    cdr = object.__new__(CarbonDioxideRemoval)
    dataset = solvent_dataset()
    dataset["exchanges"].pop()
    with pytest.raises(ValueError, match="exactly one tagged"):
        cdr.adjust_dac_water(
            dataset, "direct air capture (solvent, high-temp)", scaling_factor=1.1
        )


def test_energy_efficiency_adjustment_does_not_change_dac_water():
    cdr = object.__new__(CarbonDioxideRemoval)
    cdr.year = 2050
    cdr._get_cdr_efficiency = lambda technology, region, carrier: 2
    dataset = solvent_dataset()
    dataset["exchanges"].extend(
        [
            {
                "name": "market group for electricity, medium voltage",
                "product": "electricity, medium voltage",
                "amount": 10,
                "unit": "kilowatt hour",
                "type": "technosphere",
            },
            {
                "name": "market for heat, district or industrial, natural gas",
                "product": "heat, district or industrial, natural gas",
                "amount": 20,
                "unit": "megajoule",
                "type": "technosphere",
            },
        ]
    )
    technology = "direct air capture (solvent, high-temp)"
    cdr.adjust_dac_water(dataset, technology, scaling_factor=1.2)
    water_before = [exchange["amount"] for exchange in dataset["exchanges"][:2]]

    cdr.adjust_cdr_efficiency(dataset, technology)

    assert [exchange["amount"] for exchange in dataset["exchanges"][:2]] == (
        pytest.approx(water_before)
    )
    assert dataset["log parameters"]["DAC operational water scaling factor"] == (
        pytest.approx(1.2)
    )
