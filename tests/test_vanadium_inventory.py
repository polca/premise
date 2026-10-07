"""Source-inventory regressions for issue #286 (no licensed database required)."""

import copy
import math

import numpy as np
import pytest
from bw2io import ExcelImporter
from bw2io.strategies import csv_restore_tuples
from openpyxl import load_workbook

from premise.filesystem_constants import INVENTORY_DIR
from premise.inventory_imports import DefaultInventory
from premise.metals import correct_metal_resource_exchanges

PATH = INVENTORY_DIR / "lci-batteries-vanadium.xlsx"


@pytest.fixture(scope="module")
def inventory():
    return csv_restore_tuples(ExcelImporter(PATH).data)


def activity(inventory, name, location="ZA"):
    matches = [
        ds for ds in inventory if ds["name"] == name and ds["location"] == location
    ]
    assert len(matches) == 1
    return matches[0]


def amount(dataset, name):
    matches = [e for e in dataset["exchanges"] if e["name"] == name]
    assert len(matches) == 1
    return matches[0]["amount"]


def test_mining_allocation_preserves_element_specific_extraction(inventory):
    mine = activity(inventory, "vanadium bearing magnetite production")
    # Joint production from Weber S9; matching mine-product price basis.
    magnetite_value = 1.53 * 0.075168516669567
    ilmenite_value = 0.46 * 0.176011
    share = magnetite_value / (magnetite_value + ilmenite_value)
    assert amount(mine, "market for electricity, medium voltage") == pytest.approx(
        0.0144 / 1.53 * share
    )
    # Complementary product burdens reconstruct the unallocated batch.
    allocated = amount(mine, "market for electricity, medium voltage") * 1.53
    assert allocated + 0.0144 * (1 - share) == pytest.approx(0.0144)
    assert amount(mine, "Iron") == pytest.approx(1.08 / 1.53)
    assert amount(mine, "Vanadium") == pytest.approx(0.019 / 1.53)
    assert not any(
        e["name"] in ("Titanium", "ilmenite - magnetite mine operation")
        for e in mine["exchanges"]
    )


def test_refinery_allocates_coproduct_and_keeps_gross_chemical_input(inventory):
    refinery = activity(inventory, "vanadium pentoxide production")
    # 2010 USD prices; a short ton contains 2000 lb. Both products have 1 kg output.
    share = 6.46 / (6.46 + 140 / 2000)
    assert amount(refinery, "market for sodium sulfate, anhydrite") == pytest.approx(
        0.5 * share
    )
    assert amount(refinery, "vanadium slag production") == pytest.approx(1.35 * share)
    assert amount(refinery, "Carbon dioxide, fossil") == pytest.approx(0.16 * share)
    # Recyclable outputs and waste treatment retain their direction.
    assert amount(refinery, "market for iron scrap, unsorted") == pytest.approx(
        -0.065 * share
    )
    waste = "treatment of spent solvent mixture, hazardous waste incineration, with energy recovery"
    assert amount(refinery, waste) == pytest.approx(-0.0632 * share)


def test_intermediate_normalization_does_not_allocate_slag_twice(inventory):
    slag = activity(inventory, "vanadium slag production")
    assert amount(
        slag, "vanadium pentoxide bearing cast iron, production"
    ) == pytest.approx(1.32 * 0.5 / 0.0613)
    assert amount(slag, "market for electricity, medium voltage") == pytest.approx(
        0.54 / 0.1226
    )
    iron = activity(inventory, "vanadium pentoxide bearing cast iron, production")
    assert amount(
        iron, "pre-reduced vanadium pentoxide bearing magnetite production"
    ) == pytest.approx(1.46 / 1.32)
    reduced = activity(
        inventory, "pre-reduced vanadium pentoxide bearing magnetite production"
    )
    assert amount(reduced, "vanadium bearing magnetite production") == pytest.approx(
        1.53 / 1.46
    )
    for ds in inventory:
        if ds["location"] == "ZA":
            production = [e for e in ds["exchanges"] if e["type"] == "production"]
            assert len(production) == 1
            assert production[0]["amount"] == 1


def test_chinese_allocation_is_preserved(inventory):
    refinery = activity(inventory, "vanadium pentoxide production", "CN")
    assert amount(
        refinery, "vanadium pentoxide and unalloyed steel production"
    ) == pytest.approx(0.367078024709409 / 0.01926525)


def test_workbook_formula_caches_survive_editing():
    formulas = load_workbook(PATH, data_only=False)
    values = load_workbook(PATH, data_only=True)
    for sheet in formulas:
        for row in sheet:
            for cell in row:
                if cell.data_type == "f":
                    assert (
                        values[sheet.title][cell.coordinate].value is not None
                    ), cell.coordinate


@pytest.mark.parametrize("version", ["3.9", "3.10", "3.11", "3.12"])
def test_migration_preserves_allocation_and_uncertainty(version, monkeypatch):
    monkeypatch.setattr(
        DefaultInventory, "display_unlinked_exchanges", lambda self: None
    )
    importer = DefaultInventory([], "3.9", version, PATH, "cutoff", True)
    importer.prepare_inventory()
    data = importer.import_db.data
    assert len(data) == 13
    mine = activity(data, "vanadium bearing magnetite production")
    assert amount(mine, "Vanadium") == pytest.approx(0.019 / 1.53)
    assert not any(
        e["name"] == "ilmenite - magnetite mine operation" for e in mine["exchanges"]
    )
    for ds in data:
        if ds["location"] != "ZA":
            continue
        for exchange in ds["exchanges"]:
            assert math.isfinite(exchange["amount"])
            if exchange.get("uncertainty type") == 2:
                assert math.exp(exchange["loc"]) == pytest.approx(
                    abs(exchange["amount"])
                )
                assert bool(exchange.get("negative", False)) == (exchange["amount"] < 0)


@pytest.mark.parametrize("apply_metals", [False, True])
def test_titanium_credit_regression_through_resource_correction(
    inventory, apply_metals
):
    mine = copy.deepcopy(activity(inventory, "vanadium bearing magnetite production"))
    # A one-input background isolates the avoided-ilmenite mechanism in #286.
    # Matrix columns are activities, rows are products; mine demand is 1 kg.
    old_credit = -0.46 / 1.53
    old_matrix = np.array([[1.0, 0.0], [-old_credit, 1.0]])
    old_supply = np.linalg.solve(old_matrix, [1.0, 0.0])
    assert old_supply[1] * 0.324 < 0
    if apply_metals:
        assert correct_metal_resource_exchanges(mine, strict=True)
    credit = sum(
        e["amount"]
        for e in mine["exchanges"]
        if e["name"] == "ilmenite - magnetite mine operation"
    )
    supply = np.linalg.solve(np.array([[1.0, 0.0], [-credit, 1.0]]), [1.0, 0.0])
    direct_titanium = sum(
        e["amount"] for e in mine["exchanges"] if e["name"] == "Titanium"
    )
    assert direct_titanium + supply[1] * 0.324 == 0
    assert amount(mine, "Vanadium") == pytest.approx(0.019 / 1.53)


def test_vanadium_content_changes_invalidate_only_inventory_cache(
    tmp_path, monkeypatch
):
    import premise.new_database as module

    source = tmp_path / "vanadium.xlsx"
    source.write_bytes(PATH.read_bytes())
    monkeypatch.setattr(module, "FILEPATH_VANADIUM", source)
    obj = object.__new__(module.NewDatabase)
    obj.keep_imports_uncertainty = False
    obj.keep_source_db_uncertainty = False
    old_inventory = obj._database_cache_path("source", inventories=True)
    old_background = obj._database_cache_path("source")
    source.write_bytes(source.read_bytes() + b"different workbook content")
    assert obj._database_cache_path("source", inventories=True) != old_inventory
    assert obj._database_cache_path("source") == old_background
