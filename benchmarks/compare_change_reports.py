"""Compare benchmark outputs without loading the detailed audit into memory."""

from __future__ import annotations

import argparse
from itertools import zip_longest
import json
from pathlib import Path

import openpyxl
import pyarrow.parquet as pq


def compare(left: Path, right: Path) -> dict:
    left_metrics = json.loads((left / "metrics.json").read_text())
    right_metrics = json.loads((right / "metrics.json").read_text())
    old = pq.ParquetFile(left_metrics["details"])
    new = pq.ParquetFile(right_metrics["details"])
    assert old.schema_arrow == new.schema_arrow, "Parquet schema mismatch"
    assert old.metadata.num_rows == new.metadata.num_rows, "Audit row count mismatch"
    columns = [name for name in old.schema_arrow.names if name != "report_id"]
    rows = 0
    for a, b in zip_longest(
        old.iter_batches(batch_size=4096, columns=columns),
        new.iter_batches(batch_size=4096, columns=columns),
    ):
        assert a is not None and b is not None, "Audit batch count mismatch"
        assert a.equals(b), f"Audit differs in batch beginning at row {rows}"
        rows += a.num_rows

    old_book = openpyxl.load_workbook(left_metrics["workbook"])
    new_book = openpyxl.load_workbook(right_metrics["workbook"])
    try:
        assert old_book.sheetnames == new_book.sheetnames, "Workbook sheets differ"
        for name in old_book.sheetnames:
            a, b = old_book[name], new_book[name]
            assert (a.max_row, a.max_column) == (b.max_row, b.max_column), name
            assert a.freeze_panes == b.freeze_panes, name
            assert [table.ref for table in a.tables.values()] == [
                table.ref for table in b.tables.values()
            ], name
            for old_row, new_row in zip(a.iter_rows(), b.iter_rows()):
                for x, y in zip(old_row, new_row):
                    # These three Overview values intentionally identify a run.
                    volatile = (
                        name == "Overview"
                        and x.column == 2
                        and a.cell(x.row, 1).value
                        in {"Report ID", "Generated (UTC)", "Detailed audit"}
                    )
                    if not volatile:
                        assert (
                            x.value == y.value
                        ), f"{name}!{x.coordinate}: cell mismatch"
                    assert (
                        x.style_id == y.style_id
                    ), f"{name}!{x.coordinate}: style mismatch"
    finally:
        old_book.close()
        new_book.close()
    return dict(
        equal=True,
        rows=rows,
        speedup=left_metrics["seconds"] / right_metrics["seconds"],
        peak_rss_ratio=right_metrics["peak_rss_bytes"] / left_metrics["peak_rss_bytes"],
    )


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("baseline", type=Path)
    parser.add_argument("candidate", type=Path)
    args = parser.parse_args()
    print(json.dumps(compare(args.baseline, args.candidate), indent=2))
