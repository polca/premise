"""Compare trusted local captures from profile_update_workflow.py exactly."""

from __future__ import annotations

import argparse
from copy import deepcopy
from itertools import islice, zip_longest
import json
from pathlib import Path
import pickle
import re
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from benchmarks.compare_build_outputs import canonicalize


def normalized_metadata(metadata, *, update_only=False):
    metadata = deepcopy(metadata)
    metadata.pop("build_id", None)
    report = metadata.get("validation")
    if report:
        phases = report.get("phase_results", [])
        if update_only:
            phases = [phase for phase in phases if phase["kind"] != "export"]
            report["phase_results"] = phases
        for phase in phases:
            phase.pop("elapsed_seconds", None)
    provenance = metadata.get("provenance")
    if provenance:
        provenance.pop("build_id", None)
        for event in provenance.get("events", []):
            # Full semantic activity identity remains; generated codes vary.
            activity = event.get("activity")
            if activity:
                activity.pop("code", None)
    return canonicalize(metadata)


def differences(left, right, path="metadata"):
    if type(left) is not type(right):
        yield path, left, right
    elif isinstance(left, dict):
        for key in sorted(left.keys() | right.keys()):
            if key not in left or key not in right:
                yield f"{path}.{key}", left.get(key), right.get(key)
            else:
                yield from differences(left[key], right[key], f"{path}.{key}")
    elif isinstance(left, list):
        if len(left) != len(right):
            yield f"{path}.length", len(left), len(right)
        for index, (a, b) in enumerate(zip(left, right)):
            yield from differences(a, b, f"{path}[{index}]")
    elif left != right:
        yield path, left, right


def compare_exports(left, right):
    import openpyxl
    import pyarrow.parquet as pq

    files = [
        {path.name: path for path in (root / "simapro").glob("*.csv")}
        for root in (left, right)
    ]
    assert files[0] and files[0].keys() == files[1].keys(), "SimaPro files differ"
    identity_pairs = 0
    for name in files[0]:
        identity_pairs += compare_simapro(files[0][name], files[1][name])
        print(f"Matched SimaPro CSV: {name}", file=sys.stderr, flush=True)

    def one(root, pattern):
        matches = list((root / "reports").glob(pattern))
        assert len(matches) == 1, (root, pattern, matches)
        return matches[0]

    audits = [
        pq.ParquetFile(one(root, "change-report-*.parquet")) for root in (left, right)
    ]
    assert audits[0].schema_arrow == audits[1].schema_arrow
    assert audits[0].metadata.num_rows == audits[1].metadata.num_rows
    run_fields = {"report_id", "build_id"}
    columns = [
        name
        for name in audits[0].schema_arrow.names
        if name not in run_fields | {"activity_code"}
    ]
    run_ids = [{}, {}]
    activity_codes, reverse_activity_codes = {}, {}
    rows = 0
    for a, b in zip_longest(*(audit.iter_batches(batch_size=4096) for audit in audits)):
        assert a is not None and b is not None
        for index, batch in enumerate((a, b)):
            for name in run_fields:
                unique = batch.column(batch.schema.get_field_index(name)).unique()
                assert len(unique) == 1, (name, rows, "inconsistent run identifier")
                value = unique[0].as_py()
                assert isinstance(value, str) and value
                assert run_ids[index].setdefault(name, value) == value
        for old_code, new_code in zip(
            a.column("activity_code").to_pylist(), b.column("activity_code").to_pylist()
        ):
            if old_code is None or new_code is None:
                assert old_code is new_code, (rows, "activity code presence differs")
            else:
                assert activity_codes.setdefault(old_code, new_code) == new_code
                assert reverse_activity_codes.setdefault(new_code, old_code) == old_code
        assert a.select(columns).equals(
            b.select(columns)
        ), f"Change report differs at row {rows}"
        rows += a.num_rows
    print(f"Matched detailed audit: {rows} rows", file=sys.stderr, flush=True)

    for pattern in ("change-report-*.xlsx", "scenario_report_*.xlsx"):
        books = [openpyxl.load_workbook(one(root, pattern)) for root in (left, right)]
        try:
            assert books[0].sheetnames == books[1].sheetnames, pattern
            for name in books[0].sheetnames:
                a, b = [book[name] for book in books]
                assert (a.max_row, a.max_column) == (b.max_row, b.max_column), (
                    pattern,
                    name,
                )
                assert a.freeze_panes == b.freeze_panes
                assert [table.ref for table in a.tables.values()] == [
                    table.ref for table in b.tables.values()
                ]
                for row_a, row_b in zip(a.iter_rows(), b.iter_rows()):
                    for x, y in zip(row_a, row_b):
                        volatile = (
                            pattern == "change-report-*.xlsx"
                            and name == "Overview"
                            and x.column == 2
                            and a.cell(x.row, 1).value
                            in {"Report ID", "Generated (UTC)", "Detailed audit"}
                        )
                        if not volatile:
                            assert x.value == y.value, (
                                pattern,
                                name,
                                x.coordinate,
                                x.value,
                                y.value,
                            )
                        assert x.style_id == y.style_id, (pattern, name, x.coordinate)
        finally:
            for book in books:
                book.close()
        print(f"Matched workbook: {pattern}", file=sys.stderr, flush=True)
    return {
        "csv_files_equal": len(files[0]),
        "csv_identity_pairs_checked": identity_pairs,
        "change_report_rows_equal": rows,
        "report_activity_code_pairs_checked": len(activity_codes),
        "workbooks_equal": 2,
    }


def compare_simapro(left, right):
    """Compare every byte except a consistent bijection of generated IDs."""
    ids = re.compile(rb"((?:^|\| )ID(?:: | = ))([0-9a-f]{32})")
    forward, reverse = {}, {}
    with left.open("rb") as a, right.open("rb") as b:
        for line, (x, y) in enumerate(zip_longest(a, b), 1):
            assert x is not None and y is not None, (left.name, "line count")
            assert ids.sub(rb"\1<id>", x) == ids.sub(rb"\1<id>", y), (left.name, line)
            for old, new in zip(ids.finditer(x), ids.finditer(y)):
                old_id, new_id = old[2], new[2]
                assert forward.setdefault(old_id, new_id) == new_id, (
                    left.name,
                    line,
                    "inconsistent ID mapping",
                )
                assert reverse.setdefault(new_id, old_id) == old_id, (
                    left.name,
                    line,
                    "non-unique ID mapping",
                )
    return len(forward)


def compare(left, right, *, update_only=False, allow_subset=False):
    roots = [Path(left), Path(right)]
    manifests = [
        json.loads((root / "snapshots/manifest.json").read_text()) for root in roots
    ]
    indexed = [
        {json.dumps(row["identity"]): (index, row) for index, row in enumerate(rows)}
        for rows in manifests
    ]
    if not allow_subset:
        assert indexed[0].keys() == indexed[1].keys(), "Scenario identities differ"
    results = []
    for identity, (candidate_index, candidate) in indexed[1].items():
        baseline_index, baseline = indexed[0][identity]
        assert baseline == candidate, f"Inventory values/order differ: {identity}"
        metadata = [
            normalized_metadata(
                pickle.loads(
                    (root / f"snapshots/scenario-{index}-metadata.pickle").read_bytes()
                ),
                update_only=update_only,
            )
            for root, index in zip(roots, (baseline_index, candidate_index))
        ]
        changes = list(islice(differences(*metadata), 10))
        assert not changes, json.dumps(changes, indent=2)
        results.append({**candidate, "metadata_equal": True})
        print(
            f"Matched inventory and metadata: {identity}", file=sys.stderr, flush=True
        )
    result = {"equal": True, "scenarios": results}
    if not update_only:
        result["exports"] = compare_exports(*roots)
    return result


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("baseline", type=Path)
    parser.add_argument("candidate", type=Path)
    parser.add_argument("--update-only", action="store_true")
    parser.add_argument("--allow-subset", action="store_true")
    args = parser.parse_args()
    print(
        json.dumps(
            compare(
                args.baseline,
                args.candidate,
                update_only=args.update_only,
                allow_subset=args.allow_subset,
            ),
            indent=2,
        )
    )
